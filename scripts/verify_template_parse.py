#!/usr/bin/env python3
"""Verify that generated templates parse correctly with FileParser

Generates in-memory CSV and XLSX templates using internal field names
and feeds them to FileParser.parse to ensure no missing-field errors.
"""
import sys
import io
import traceback

ROOT = r"c:\ECOLAB-ETS\RnD\nucleIQ"
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

def main():
    try:
        from backend.data_management.templates import get_template
        from backend.data_management.utils import FileParser, OPENPYXL_AVAILABLE
    except Exception:
        print('Failed to import project modules:')
        traceback.print_exc()
        sys.exit(2)

    modules = ['classes', 'staff', 'students']

    for module in modules:
        print(f'\nTesting module: {module}')
        template = get_template(module)
        if not template:
            print('  No template found; skipping')
            continue

        # CSV test
        headers = [f.name for f in template.fields]
        sample = [str(f.sample_value or '') for f in template.fields]
        csv_content = ','.join(headers) + '\n' + ','.join(sample) + '\n'
        try:
            data, _, file_type = FileParser.parse(io.BytesIO(csv_content.encode('utf-8')), filename=f'{module}.csv')
            print(f'  CSV parse: {len(data)} rows, keys={list(data[0].keys()) if data else []}')
        except Exception as e:
            print('  CSV parse failed:')
            traceback.print_exc()

        # XLSX test (if available)
        if OPENPYXL_AVAILABLE:
            try:
                import openpyxl
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = template.display_name or module
                # header row - internal names
                for col, name in enumerate(headers, start=1):
                    ws.cell(row=1, column=col, value=name)
                # sample row
                for col, val in enumerate(sample, start=1):
                    ws.cell(row=2, column=col, value=val)

                bio = io.BytesIO()
                wb.save(bio)
                bio.seek(0)

                data, _, file_type = FileParser.parse(bio, filename=f'{module}.xlsx')
                print(f'  XLSX parse: {len(data)} rows, keys={list(data[0].keys()) if data else []}')
            except Exception:
                print('  XLSX parse failed:')
                traceback.print_exc()
        else:
            print('  openpyxl not available; skipping XLSX test')

    print('\nVerify complete')

if __name__ == '__main__':
    main()
