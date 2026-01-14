"""
Data Management Views - Import, Export, Backup functionality
Tenant Admin Only Access
"""
import io
import csv
import json
import zipfile
from datetime import datetime
from django.http import HttpResponse
from django.db import transaction
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from core.permissions import IsTenantAdmin

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    import openpyxl.utils
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from students.models import Student, StudentEnrollment
from staff.models import Staff
from tenants.models import GradeLevel, Section, Subject, AcademicYear
from fees.models import FeeCategory, FeeStructure, FeeAllocation


def parse_date_flexible(date_str):
    """
    Parse date string in multiple formats:
    - dd-mm-yyyy
    - dd/mm/yyyy  
    - yyyy-mm-dd (ISO format)
    Returns a date object or None if parsing fails.
    """
    if not date_str:
        return None
    
    # If already a date object, return as is
    if hasattr(date_str, 'year'):
        return date_str
    
    date_str = str(date_str).strip()
    
    # Try different formats
    formats = [
        '%d-%m-%Y',  # dd-mm-yyyy
        '%d/%m/%Y',  # dd/mm/yyyy
        '%Y-%m-%d',  # yyyy-mm-dd (ISO)
        '%d-%m-%y',  # dd-mm-yy
        '%d/%m/%y',  # dd/mm/yy
    ]
    
    for fmt in formats:
        try:
            return datetime.strptime(date_str, fmt).date()
        except ValueError:
            continue
    
    # If all formats fail, return None
    return None


class TemplateInfo:
    """Template configuration for each module"""
    
    TEMPLATES = {
        'students': {
            'name': 'Students',
            'model': Student,
            'required_fields': ['first_name', 'last_name', 'admission_number', 'date_of_birth', 'gender', 'father_name', 'mother_name', 'father_phone', 'mother_phone'],
            'optional_fields': ['email', 'phone', 'address', 'blood_group', 
                              'aadhar_number', 'nationality', 'religion', 'caste', 'admission_date', 
                              'class_name', 'section_name'],
            'unique_field': 'admission_number',
            'sample_data': {
                'first_name': 'John',
                'last_name': 'Doe',
                'admission_number': 'STU001',
                'date_of_birth': '15-05-2010',
                'gender': 'M',
                'address': '123 Main Street, City',
                'father_name': 'Robert Doe',
                'mother_name': 'Mary Doe',
                'father_phone': '9876543210',
                'mother_phone': '9876543211',
                'class_name': 'Class 10',
                'section_name': 'A',
                'admission_date': '01-04-2024'
            }
        },
        'staff': {
            'name': 'Staff',
            'model': Staff,
            'required_fields': ['first_name', 'last_name', 'employee_id', 'email', 'date_of_joining'],
            'optional_fields': ['phone', 'address', 'blood_group', 'date_of_birth', 'gender', 
                              'qualification', 'experience_years', 'department', 'designation'],
            'unique_field': 'employee_id',
            'sample_data': {
                'first_name': 'Jane',
                'last_name': 'Smith',
                'employee_id': 'EMP001',
                'email': 'jane@school.com',
                'date_of_joining': '01-06-2020',
                'phone': '9876543212',
                'gender': 'Female',
                'department': 'Mathematics',
                'designation': 'Senior Teacher',
                'date_of_birth': '15-03-1985'
            }
        },
        'classes': {
            'name': 'Classes & Sections',
            'model': GradeLevel,
            'required_fields': ['class_name', 'section_name'],
            'optional_fields': ['room_number', 'capacity'],
            'unique_field': None,  # Composite unique
            'sample_data': {
                'class_name': 'Class 10',
                'section_name': 'A',
                'room_number': '101',
                'capacity': '40'
            }
        },
        'fee_structures': {
            'name': 'Fee Structures',
            'model': FeeStructure,
            'required_fields': ['fee_type', 'class_name', 'amount', 'frequency'],
            'optional_fields': ['due_day', 'description', 'is_mandatory', 'late_fee_percent'],
            'unique_field': None,
            'sample_data': {
                'fee_type': 'Tuition Fee',
                'class_name': 'Class 10',
                'amount': '5000',
                'frequency': 'monthly',
                'due_day': '10',
                'is_mandatory': 'true'
            }
        },
        'subjects': {
            'name': 'Subjects',
            'model': Subject,
            'required_fields': ['subject_name', 'subject_code', 'class_name'],
            'optional_fields': ['credit_hours', 'is_elective', 'max_marks'],
            'unique_field': 'subject_code',
            'sample_data': {
                'subject_name': 'Mathematics',
                'subject_code': 'MATH10',
                'class_name': 'Class 10',
                'credit_hours': '5',
                'is_elective': 'false'
            }
        },
        'fee_allocations': {
            'name': 'Fee Allocations',
            'model': FeeAllocation,
            'required_fields': ['admission_number', 'fee_type', 'amount'],
            'optional_fields': ['discount_percent', 'discount_reason'],
            'unique_field': None,
            'sample_data': {
                'admission_number': 'STU001',
                'fee_type': 'Tuition Fee',
                'amount': '5000',
                'discount_percent': '10',
                'discount_reason': 'Sibling Discount'
            }
        },
        'parents': {
            'name': 'Parents',
            'model': Student,
            'required_fields': ['admission_number', 'father_name', 'mother_name', 'father_phone'],
            'optional_fields': ['father_email', 'mother_phone', 'mother_email'],
            'unique_field': 'admission_number',
            'sample_data': {
                'admission_number': 'STU001',
                'father_name': 'Robert Doe',
                'mother_name': 'Mary Doe',
                'father_phone': '9876543211'
            }
        },
        'student_photos': {
            'name': 'Student Photos',
            'model': Student,
            'required_fields': ['admission_number_as_filename'],
            'optional_fields': ['format_jpg_png'],
            'unique_field': None,
            'sample_data': {
                'admission_number_as_filename': 'The ZIP should contain images named like STU001.jpg, STU002.png etc.'
            }
        }
    }


class DownloadTemplateView(APIView):
    """Download Excel template for a specific module"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request, module):
        if module not in TemplateInfo.TEMPLATES:
            return Response({'error': 'Invalid module'}, status=400)
        
        template_config = TemplateInfo.TEMPLATES[module]
        
        if EXCEL_AVAILABLE:
            return self._generate_excel_template(module, template_config)
        else:
            return self._generate_csv_template(module, template_config)
    
    def _generate_excel_template(self, module, config):
        wb = Workbook()
        ws = wb.active
        ws.title = config['name']
        
        # Header styling
        header_font = Font(bold=True, color='FFFFFF')
        required_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
        optional_fill = PatternFill(start_color='64748B', end_color='64748B', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Write headers
        all_fields = config['required_fields'] + config['optional_fields']
        for col, field in enumerate(all_fields, 1):
            cell = ws.cell(row=1, column=col, value=field)
            cell.font = header_font
            cell.fill = required_fill if field in config['required_fields'] else optional_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max(15, len(field) + 2)
        
        # Write sample data row
        sample = config['sample_data']
        for col, field in enumerate(all_fields, 1):
            cell = ws.cell(row=2, column=col, value=sample.get(field, ''))
            cell.border = thin_border
        
        # Add instructions sheet
        instructions_ws = wb.create_sheet('Instructions')
        instructions = [
            ('📋 IMPORT INSTRUCTIONS', '', ''),
            ('', '', ''),
            ('📅 DATE FORMAT:', '', ''),
            ('All date fields accept BOTH formats:', '', ''),
            ('   • dd-mm-yyyy (e.g., 15-05-2010)', '', ''),
            ('   • dd/mm/yyyy (e.g., 15/05/2010)', '', ''),
            ('', '', ''),
            ('📝 FIELD DETAILS:', '', ''),
            ('Field', 'Required', 'Description'),
        ]
        for field in config['required_fields']:
            description = f'Required field'
            if 'date' in field.lower():
                description = f'Required date field. Format: dd-mm-yyyy or dd/mm/yyyy'
            instructions.append((field, 'Yes', description))
        for field in config['optional_fields']:
            description = f'Optional field'
            if 'date' in field.lower():
                description = f'Optional date field. Format: dd-mm-yyyy or dd/mm/yyyy'
            instructions.append((field, 'No', description))
        
        # Add additional instructions
        instructions.append(('', '', ''))
        instructions.append(('💡 TIPS:', '', ''))
        instructions.append(('• Blue headers are REQUIRED fields', '', ''))
        instructions.append(('• Gray headers are OPTIONAL fields', '', ''))
        instructions.append(('• Do not modify the header row', '', ''))
        instructions.append(('• Sample data is provided in row 2 of the main sheet', '', ''))
        
        for row_num, row_data in enumerate(instructions, 1):
            for col_num, value in enumerate(row_data, 1):
                cell = instructions_ws.cell(row=row_num, column=col_num, value=value)
                if row_num == 1 or row_num == 9:
                    cell.font = Font(bold=True, size=12)
                elif '📅' in str(value) or '📝' in str(value) or '💡' in str(value):
                    cell.font = Font(bold=True)
        
        # Adjust column widths for instructions
        instructions_ws.column_dimensions['A'].width = 45
        instructions_ws.column_dimensions['B'].width = 12
        instructions_ws.column_dimensions['C'].width = 55
        
        # Return response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{module}_template.xlsx"'
        wb.save(response)
        return response
    
    def _generate_csv_template(self, module, config):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{module}_template.csv"'
        
        writer = csv.writer(response)
        all_fields = config['required_fields'] + config['optional_fields']
        writer.writerow(all_fields)
        
        # Sample row
        sample = config['sample_data']
        writer.writerow([sample.get(field, '') for field in all_fields])
        
        return response


class ImportDataView(APIView):
    """Import data from Excel/CSV with duplicate detection"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        module = request.data.get('module')
        file = request.FILES.get('file')
        skip_duplicates = request.data.get('skip_duplicates', 'true') == 'true'
        
        if not module or module not in TemplateInfo.TEMPLATES:
            return Response({'error': 'Invalid module'}, status=400)
        
        if not file:
            return Response({'error': 'No file provided'}, status=400)
        
        config = TemplateInfo.TEMPLATES[module]
        tenant = request.user.tenant
        
        try:
            # Handle ZIP photo upload
            if module == 'student_photos' and file.name.endswith('.zip'):
                return Response(self._process_photo_zip(file, tenant))

            # Parse file
            if file.name.endswith('.csv'):
                data = self._parse_csv(file)
            elif file.name.endswith(('.xlsx', '.xls')):
                if not EXCEL_AVAILABLE:
                    return Response({'error': 'Excel support not available. Please use CSV.'}, status=400)
                data = self._parse_excel(file)
            else:
                return Response({'error': 'Unsupported file format. Use .xlsx or .csv'}, status=400)
            
            if not data:
                return Response({'error': 'No data found in file'}, status=400)
            
            # Process data
            result = self._process_import(module, config, data, tenant, skip_duplicates)
            return Response(result)
            
        except Exception as e:
            return Response({'error': str(e)}, status=500)
    
    def _parse_csv(self, file):
        content = file.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(content))
        return list(reader)

    def _process_photo_zip(self, file, tenant):
        """Process a ZIP file containing student photos named by admission number"""
        success = 0
        failed = 0
        errors = []
        
        try:
            with zipfile.ZipFile(file) as z:
                for filename in z.namelist():
                    # Skip directories and non-image files
                    if filename.endswith('/') or not filename.lower().endswith(('.jpg', '.jpeg', '.png')):
                        continue
                    
                    # Extract admission number from filename (e.g. STU001.jpg -> STU001)
                    name_part = filename.split('/')[-1] # Handle nested files
                    admission_number = name_part.rsplit('.', 1)[0]
                    
                    try:
                        student = Student.objects.filter(tenant=tenant, admission_number=admission_number).first()
                        if student:
                            # Read file content
                            photo_data = z.read(filename)
                            # Save to student photo field
                            from django.core.files.base import ContentFile
                            student.photo.save(f"{admission_number}.{name_part.rsplit('.', 1)[1]}", ContentFile(photo_data), save=True)
                            success += 1
                        else:
                            failed += 1
                            errors.append(f"Student {admission_number} not found (File: {filename})")
                    except Exception as e:
                        failed += 1
                        errors.append(f"Error processing {filename}: {str(e)}")
        except Exception as e:
            return {'success': 0, 'failed': 0, 'duplicates': [], 'errors': [f"Invalid ZIP: {str(e)}"]}
            
        return {
            'success': success,
            'failed': failed,
            'duplicates': [],
            'errors': errors[:20]
        }
    
    def _parse_excel(self, file):
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1] if cell.value]
        data = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):  # Skip empty rows
                row_dict = {}
                for i, header in enumerate(headers):
                    if i < len(row):
                        value = row[i]
                        if value is not None:
                            row_dict[header] = str(value) if not isinstance(value, str) else value
                if row_dict:
                    data.append(row_dict)
        
        return data
    
    def _process_import(self, module, config, data, tenant, skip_duplicates):
        success = 0
        failed = 0
        duplicates = []
        errors = []
        
        unique_field = config.get('unique_field')
        model = config['model']
        required_fields = config['required_fields']
        
        for row_num, row in enumerate(data, start=2):
            try:
                # Validate required fields
                missing = [f for f in required_fields if not row.get(f)]
                if missing:
                    errors.append(f"Row {row_num}: Missing required fields: {', '.join(missing)}")
                    failed += 1
                    continue
                
                # Check for duplicates
                if unique_field and row.get(unique_field):
                    existing = self._find_duplicate(model, unique_field, row[unique_field], tenant)
                    if existing:
                        if skip_duplicates:
                            duplicates.append({
                                'row': row_num,
                                'data': row,
                                'existingRecord': self._model_to_dict(existing),
                                'field': unique_field,
                                'value': row[unique_field]
                            })
                            continue
                
                # Create record
                self._create_record(module, model, row, tenant)
                success += 1
                
            except Exception as e:
                errors.append(f"Row {row_num}: {str(e)}")
                failed += 1
        
        return {
            'success': success,
            'failed': failed,
            'duplicates': duplicates,
            'errors': errors[:20]  # Limit errors returned
        }
    
    def _find_duplicate(self, model, field, value, tenant):
        try:
            filter_kwargs = {field: value, 'tenant': tenant}
            return model.objects.filter(**filter_kwargs).first()
        except:
            return None
    
    def _model_to_dict(self, instance):
        """Convert model instance to dictionary for comparison"""
        if not instance:
            return {}
        
        result = {}
        for field in instance._meta.fields:
            value = getattr(instance, field.name, None)
            if value is not None:
                if hasattr(value, 'isoformat'):
                    value = value.isoformat()
                result[field.name] = str(value) if not isinstance(value, (str, int, float, bool)) else value
        return result
    
    def _create_record(self, module, model, row, tenant):
        """Create a new record based on module type"""
        if module == 'students':
            self._create_student(row, tenant)
        elif module == 'staff':
            self._create_staff(row, tenant)
        elif module == 'classes':
            self._create_class_section(row, tenant)
        elif module == 'fee_structures':
            self._create_fee_structure(row, tenant)
        elif module == 'subjects':
            self._create_subject(row, tenant)
        elif module == 'fee_allocations':
            self._create_fee_allocation(row, tenant)
        elif module == 'parents':
            self._create_parent(row, tenant)
        else:
            raise ValueError(f"Unknown module: {module}")
    
    @transaction.atomic
    def _create_parent(self, row, tenant):
        admission_number = row['admission_number']
        student = Student.objects.filter(tenant=tenant, admission_number=admission_number).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")
        
        # Update parent fields on student
        student.father_name = row.get('father_name', student.father_name)
        student.mother_name = row.get('mother_name', student.mother_name)
        student.father_phone = row.get('father_phone', student.father_phone)
        student.mother_phone = row.get('mother_phone', student.mother_phone)
        student.save()

    @transaction.atomic
    def _create_fee_allocation(self, row, tenant):
        admission_number = row['admission_number']
        fee_type_name = row['fee_type']
        
        student = Student.objects.filter(tenant=tenant, admission_number=admission_number).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")
            
        fee_type = FeeCategory.objects.filter(tenant=tenant, name=fee_type_name).first()
        if not fee_type:
            raise ValueError(f"Fee Category '{fee_type_name}' not found")

        # Find fee structure for this student's class and fee type
        enrollment = student.get_current_enrollment()
        if not enrollment:
            raise ValueError(f"No active enrollment found for student '{admission_number}'")
            
        academic_year = AcademicYear.objects.filter(tenant=tenant, is_active=True).first()
        if not academic_year:
            raise ValueError("No active academic year found")

        fee_structure = FeeStructure.objects.filter(
            tenant=tenant,
            academic_year=academic_year,
            class_level=enrollment.class_assigned,
            category=fee_type
        ).first()
        
        if not fee_structure:
            # Create structure if it doesn't exist? Or fail? 
            # Better create a basic one or fail. Let's fail for now to ensure user created it.
            raise ValueError(f"Fee Structure not found for {enrollment.class_assigned.name} - {fee_type_name}")
            
        FeeAllocation.objects.create(
            tenant=tenant,
            student=student,
            fee_structure=fee_structure,
            # scholarship_percentage=float(row.get('discount_percent', 0)), # Match model field
        )

    @transaction.atomic
    def _create_student(self, row, tenant):
        from users.models import User
        
        # Parse dates with flexible format support
        date_of_birth = parse_date_flexible(row.get('date_of_birth'))
        admission_date = parse_date_flexible(row.get('admission_date')) or datetime.now().date()
        
        student = Student.objects.create(
            tenant=tenant,
            first_name=row['first_name'],
            last_name=row['last_name'],
            admission_number=row['admission_number'],
            date_of_birth=date_of_birth,
            gender=row.get('gender', 'M'),
            email=row.get('email', '') or '',
            phone=row.get('phone', '') or '',
            address=row.get('address', '') or '',
            blood_group=row.get('blood_group', '') or '',
            father_name=row.get('father_name', '') or '',
            mother_name=row.get('mother_name', '') or '',
            father_phone=row.get('father_phone', '') or '',
            mother_phone=row.get('mother_phone', '') or '',
            aadhar_number=row.get('aadhar_number', '') or '',
            nationality=row.get('nationality', 'Indian') or 'Indian',
            religion=row.get('religion', '') or '',
            caste=row.get('caste', '') or '',
            admission_date=admission_date
        )
        
        # Create enrollment if class/section provided
        class_name = row.get('class_name')
        section_name = row.get('section_name')
        
        if class_name:
            try:
                class_obj = GradeLevel.objects.filter(tenant=tenant, name=class_name).first()
                if class_obj:
                    section = None
                    if section_name:
                        section = Section.objects.filter(
                            tenant=tenant, 
                            grade_level=class_obj, 
                            name=section_name
                        ).first()
                    
                    StudentEnrollment.objects.create(
                        tenant=tenant,
                        student=student,
                        class_assigned=class_obj,
                        section=section,
                        is_active=True
                    )
            except Exception as e:
                pass  # Don't fail if enrollment fails
        
        return student
    
    def _create_staff(self, row, tenant):
        # Parse dates with flexible format support
        date_of_joining = parse_date_flexible(row.get('date_of_joining'))
        date_of_birth = parse_date_flexible(row.get('date_of_birth'))
        
        staff = Staff.objects.create(
            tenant=tenant,
            first_name=row['first_name'],
            last_name=row['last_name'],
            employee_id=row['employee_id'],
            email=row['email'],
            date_of_joining=date_of_joining,
            phone=row.get('phone'),
            address=row.get('address'),
            blood_group=row.get('blood_group'),
            date_of_birth=date_of_birth,
            gender=row.get('gender'),
            qualification=row.get('qualification'),
            department=row.get('department'),
            designation=row.get('designation'),
        )
        return staff
    
    def _create_class_section(self, row, tenant):
        class_name = row['class_name']
        section_name = row['section_name']
        
        # Get or create class
        class_obj, _ = GradeLevel.objects.get_or_create(
            tenant=tenant,
            name=class_name,
            defaults={'display_order': 0}
        )
        
        # Create section
        Section.objects.get_or_create(
            tenant=tenant,
            grade_level=class_obj,
            name=section_name,
            defaults={
                'room_number': row.get('room_number'),
                'capacity': int(row.get('capacity', 40))
            }
        )
    
    def _create_fee_structure(self, row, tenant):
        fee_type_name = row['fee_type']
        class_name = row['class_name']
        
        fee_type, _ = FeeCategory.objects.get_or_create(
            tenant=tenant,
            name=fee_type_name
        )
        
        class_obj = GradeLevel.objects.filter(tenant=tenant, name=class_name).first()
        if not class_obj:
            raise ValueError(f"Class '{class_name}' not found")
        
        academic_year = AcademicYear.objects.filter(tenant=tenant, is_active=True).first()
        if not academic_year:
            raise ValueError("No active academic year found")

        FeeStructure.objects.create(
            tenant=tenant,
            academic_year=academic_year,
            category=fee_type,
            class_level=class_obj,
            amount=float(row['amount']),
            frequency=row.get('frequency', 'MONTHLY').upper(),
            due_day=int(row.get('due_day', 10)),
            description=row.get('description', ''),
            is_mandatory=row.get('is_mandatory', 'true').lower() == 'true'
        )
    
    def _create_subject(self, row, tenant):
        class_name = row['class_name']
        class_obj = GradeLevel.objects.filter(tenant=tenant, name=class_name).first()
        
        if not class_obj:
            raise ValueError(f"Class '{class_name}' not found")
        
        Subject.objects.create(
            tenant=tenant,
            name=row['subject_name'],
            code=row['subject_code'],
            # Note: Subject in tenants.models doesn't link to class directly, 
            # but ClassSubject does. Creating subject as master here.
        )


class OverrideDuplicatesView(APIView):
    """Override duplicate records with new data"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def post(self, request):
        module = request.data.get('module')
        records = request.data.get('records', [])
        
        if not module or module not in TemplateInfo.TEMPLATES:
            return Response({'error': 'Invalid module'}, status=400)
        
        config = TemplateInfo.TEMPLATES[module]
        tenant = request.user.tenant
        model = config['model']
        unique_field = config.get('unique_field')
        
        updated = 0
        errors = []
        
        for record in records:
            try:
                if unique_field:
                    value = record['data'].get(unique_field)
                    instance = model.objects.filter(tenant=tenant, **{unique_field: value}).first()
                    
                    if instance:
                        # Update fields
                        for field, val in record['data'].items():
                            if hasattr(instance, field) and val:
                                setattr(instance, field, val)
                        instance.save()
                        updated += 1
            except Exception as e:
                errors.append(f"Failed to update {record.get('value', 'unknown')}: {str(e)}")
        
        return Response({'updated': updated, 'errors': errors})


class ExportDataView(APIView):
    """Export data to Excel/CSV"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request, module):
        if module not in TemplateInfo.TEMPLATES:
            return Response({'error': 'Invalid module'}, status=400)
        
        tenant = request.user.tenant
        config = TemplateInfo.TEMPLATES[module]
        model = config['model']
        
        # Get all records for tenant
        queryset = model.objects.filter(tenant=tenant)
        
        if EXCEL_AVAILABLE:
            return self._export_excel(module, config, queryset)
        else:
            return self._export_csv(module, config, queryset)
    
    def _export_excel(self, module, config, queryset):
        wb = Workbook()
        ws = wb.active
        ws.title = config['name']
        
        # Headers
        all_fields = config['required_fields'] + config['optional_fields']
        
        header_font = Font(bold=True)
        for col, field in enumerate(all_fields, 1):
            cell = ws.cell(row=1, column=col, value=field)
            cell.font = header_font
        
        # Data rows
        for row_num, obj in enumerate(queryset, 2):
            for col, field in enumerate(all_fields, 1):
                value = getattr(obj, field, None)
                if value is not None:
                    if hasattr(value, 'isoformat'):
                        value = value.isoformat()
                    elif hasattr(value, 'name'):  # FK
                        value = value.name
                    ws.cell(row=row_num, column=col, value=str(value) if value else '')
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{module}_export_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response
    
    def _export_csv(self, module, config, queryset):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{module}_export_{datetime.now().strftime("%Y%m%d")}.csv"'
        
        writer = csv.writer(response)
        all_fields = config['required_fields'] + config['optional_fields']
        writer.writerow(all_fields)
        
        for obj in queryset:
            row = []
            for field in all_fields:
                value = getattr(obj, field, None)
                if value is not None:
                    if hasattr(value, 'isoformat'):
                        value = value.isoformat()
                    elif hasattr(value, 'name'):
                        value = value.name
                row.append(str(value) if value else '')
            writer.writerow(row)
        
        return response


class FullBackupView(APIView):
    """Create full backup of all tenant data"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request):
        tenant = request.user.tenant
        
        # Create in-memory ZIP file
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            # Export each module
            for module, config in TemplateInfo.TEMPLATES.items():
                try:
                    model = config['model']
                    queryset = model.objects.filter(tenant=tenant)
                    
                    if not queryset.exists():
                        continue
                    
                    # Create CSV content
                    csv_buffer = io.StringIO()
                    writer = csv.writer(csv_buffer)
                    
                    all_fields = config['required_fields'] + config['optional_fields']
                    writer.writerow(all_fields)
                    
                    for obj in queryset:
                        row = []
                        for field in all_fields:
                            value = getattr(obj, field, None)
                            if value is not None:
                                if hasattr(value, 'isoformat'):
                                    value = value.isoformat()
                                elif hasattr(value, 'name'):
                                    value = value.name
                            row.append(str(value) if value else '')
                        writer.writerow(row)
                    
                    # Add to ZIP
                    zip_file.writestr(f"{module}.csv", csv_buffer.getvalue())
                    
                except Exception as e:
                    # Log error but continue with other exports
                    pass
            
            # Add metadata
            metadata = {
                'tenant': str(tenant.id) if hasattr(tenant, 'id') else 'unknown',
                'backup_date': datetime.now().isoformat(),
                'modules': list(TemplateInfo.TEMPLATES.keys())
            }
            zip_file.writestr('backup_metadata.json', json.dumps(metadata, indent=2))
        
        zip_buffer.seek(0)
        
        response = HttpResponse(zip_buffer.read(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="school_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip"'
        return response
