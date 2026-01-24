"""
Data Migration Wizard - Simplified Admin Views
A user-friendly interface for data migration in Django Admin
"""

import os
import io
import csv
import json
import tempfile
from datetime import date
from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from django.db import transaction

from tenants.models import Tenant, AcademicYear
from .templates import get_template, get_all_templates
from .models import ImportJob


# =============================================================================
# TEMPLATE CONFIGURATION
# =============================================================================

# Module icons and metadata
MODULE_ICONS = {
    'students': '🎓',
    'staff': '👨‍🏫',
    'classes': '📚',
    'subjects': '📖',
    'fee_structures': '💰',
    'fee_invoices': '🧾',
    'fee_payments': '💸',
    'fee_allocations': '💳',
    'student_enrollments': '📝',
    'transport': '🚌',
    'parents': '👨‍👩‍👦',
    'user_accounts': '🔐',
    'attendance': '📋',
    'student_photos': '📷',
    'exam_results': '📝',
    'exam_schedule': '🗓️',
    'timetable': '📆',
    'library_books': '📚',
    'library_transactions': '📖',
    'payroll_payments': '💵',
    'hostel_allocations': '🏠',
    'inventory_items': '📦',
    'certificates_issued': '📜',
    'finance_journal_entries': '📒',
    'fee_discounts': '🏷️',
    'helpdesk_tickets': '🎧',
    'lms_courses': '💻',
    'lms_enrollments': '🧑‍🏫',
    'idcards': '🪪',
}

# Modules that support historical data import
HISTORICAL_MODULES = [
    'students', 'staff', 'classes', 'fee_structures', 'fee_invoices',
    'student_enrollments', 'attendance', 'fee_allocations'
]

# Academic year range options
YEAR_RANGE_OPTIONS = [
    ('1', '1 Year (Current Year Only)'),
    ('3', '3 Years'),
    ('5', '5 Years'),
    ('10', '10 Years (Full Historical Data)'),
    ('custom', 'Custom Range'),
]


def get_module_templates():
    """Get all module templates with metadata"""
    templates = get_all_templates()
    result = {}
    for name, template in templates.items():
        result[name] = {
            'template': template,
            'icon': MODULE_ICONS.get(name, '📄'),
            'description': template.description,
            'supports_historical': name in HISTORICAL_MODULES,
        }
    return result


# Create MODULE_TEMPLATES as a lazy-loaded dict
MODULE_TEMPLATES = None

def _get_templates():
    """Lazy load templates"""
    global MODULE_TEMPLATES
    if MODULE_TEMPLATES is None:
        MODULE_TEMPLATES = get_module_templates()
    return MODULE_TEMPLATES


def get_academic_years_for_range(tenant, year_count):
    """
    Generate academic year codes for the specified range.
    Returns list of academic year strings like ['2025-26', '2024-25', '2023-24']
    """
    current_year = date.today().year
    current_month = date.today().month
    
    # Determine current academic year (assuming April start)
    if current_month >= 4:
        start_year = current_year
    else:
        start_year = current_year - 1
    
    years = []
    for i in range(int(year_count)):
        year_start = start_year - i
        year_end = (year_start + 1) % 100  # Get last 2 digits
        years.append(f"{year_start}-{year_end:02d}")
    
    return years


def generate_template_with_years(module_name, year_count, include_sample_data=True):
    """
    Generate an Excel template with academic year columns.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    module_config = _get_templates().get(module_name)
    if not module_config:
        raise ValueError(f"Unknown module: {module_name}")
    
    template = module_config['template']
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = template.display_name[:31]  # Excel sheet name limit
    
    # Styles
    header_fill_required = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_fill_optional = PatternFill(start_color="6B7280", end_color="6B7280", fill_type="solid")
    header_fill_year = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    sample_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Collect all fields
    all_fields = template.fields
    required_fields = [f for f in all_fields if f.required]
    optional_fields = [f for f in all_fields if not f.required]
    
    # Write headers
    col = 1
    field_columns = {}
    
    # Required fields first
    for field in required_fields:
        cell = ws.cell(row=1, column=col, value=field.name)
        cell.fill = header_fill_required
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = max(15, len(field.name) + 2)
        field_columns[field.name] = col
        col += 1
    
    # Optional fields
    for field in optional_fields:
        cell = ws.cell(row=1, column=col, value=field.name)
        cell.fill = header_fill_optional
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = max(15, len(field.name) + 2)
        field_columns[field.name] = col
        col += 1
    
    # Add academic year column if module supports historical data
    if module_config.get('supports_historical') and int(year_count) > 1:
        cell = ws.cell(row=1, column=col, value='academic_year')
        cell.fill = header_fill_year
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 15
        field_columns['academic_year'] = col
        col += 1
    
    # Write sample data rows
    if include_sample_data:
        years = get_academic_years_for_range(None, year_count)
        
        for row_idx, year in enumerate(years[:3], start=2):  # Max 3 sample rows
            for field in all_fields:
                col_idx = field_columns.get(field.name)
                if col_idx:
                    cell = ws.cell(row=row_idx, column=col_idx, value=field.sample_value)
                    cell.fill = sample_fill
                    cell.border = thin_border
            
            # Add academic year if applicable
            if 'academic_year' in field_columns:
                cell = ws.cell(row=row_idx, column=field_columns['academic_year'], value=year)
                cell.fill = sample_fill
                cell.border = thin_border
    
    # Create Instructions sheet
    ws_inst = wb.create_sheet("Instructions")
    ws_inst.column_dimensions['A'].width = 80
    
    instructions = [
        f"📋 {template.display_name.upper()} IMPORT TEMPLATE",
        "",
        "═" * 50,
        "",
        "🔵 REQUIRED FIELDS (Blue headers - must be filled):",
        *[f"   • {f.display_name}: {f.description}" for f in required_fields],
        "",
        "⚪ OPTIONAL FIELDS (Gray headers - can be left blank):",
        *[f"   • {f.display_name}: {f.description}" for f in optional_fields],
        "",
        "═" * 50,
    ]
    
    if module_config.get('supports_historical'):
        instructions.extend([
            "",
            "🟢 ACADEMIC YEAR COLUMN (Green header):",
            f"   • Format: YYYY-YY (e.g., 2025-26)",
            f"   • You selected {year_count} year(s) of data",
            "   • Add one row per record per academic year",
            "",
        ])
    
    if template.instructions:
        instructions.extend(["", "═" * 50, ""])
        instructions.extend(template.instructions)
    
    for idx, line in enumerate(instructions, start=1):
        cell = ws_inst.cell(row=idx, column=1, value=line)
        if line.startswith("📋") or line.startswith("═"):
            cell.font = Font(bold=True, size=12)
    
    # Create Field Reference sheet
    ws_ref = wb.create_sheet("Field Reference")
    ws_ref.column_dimensions['A'].width = 25
    ws_ref.column_dimensions['B'].width = 15
    ws_ref.column_dimensions['C'].width = 40
    ws_ref.column_dimensions['D'].width = 20
    ws_ref.column_dimensions['E'].width = 10
    
    # Headers
    for col_idx, header in enumerate(['Field Name', 'Type', 'Description', 'Sample Value', 'Required'], start=1):
        cell = ws_ref.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill_required
        cell.font = header_font
        cell.border = thin_border
    
    # Field data
    for row_idx, field in enumerate(all_fields, start=2):
        ws_ref.cell(row=row_idx, column=1, value=field.display_name).border = thin_border
        ws_ref.cell(row=row_idx, column=2, value=field.field_type.value).border = thin_border
        ws_ref.cell(row=row_idx, column=3, value=field.description).border = thin_border
        ws_ref.cell(row=row_idx, column=4, value=field.sample_value).border = thin_border
        ws_ref.cell(row=row_idx, column=5, value='Yes' if field.required else 'No').border = thin_border
    
    return wb


# =============================================================================
# ADMIN VIEWS
# =============================================================================

@staff_member_required
def migration_wizard(request):
    """
    Main migration wizard view - Step 1: Select tenant and module
    """
    tenants = Tenant.objects.filter(is_active=True).order_by('name')
    
    modules = []
    for name, config in _get_templates().items():
        modules.append({
            'name': name,
            'display_name': config['template'].display_name,
            'icon': config['icon'],
            'description': config['description'],
            'supports_historical': config.get('supports_historical', False),
            'required_fields': len(config['template'].get_required_fields()),
            'optional_fields': len(config['template'].get_optional_fields()),
        })
    
    # Sort modules by common usage
    module_order = ['students', 'staff', 'classes', 'subjects', 'fee_structures', 
                    'fee_invoices', 'fee_payments', 'fee_allocations', 'fee_discounts',
                    'student_enrollments', 'attendance', 'transport', 'parents', 'user_accounts',
                    'student_photos', 'exam_results', 'exam_schedule', 'timetable',
                    'library_books', 'library_transactions', 'payroll_payments',
                    'hostel_allocations', 'inventory_items', 'certificates_issued',
                    'finance_journal_entries', 'helpdesk_tickets', 'lms_courses',
                    'lms_enrollments', 'idcards']
    modules.sort(key=lambda x: module_order.index(x['name']) if x['name'] in module_order else 999)
    
    context = {
        'title': 'Data Migration Wizard',
        'tenants': tenants,
        'modules': modules,
        'year_options': YEAR_RANGE_OPTIONS,
        'site_header': 'NucleiQ Platform Administration',
        'has_permission': True,
    }
    
    return render(request, 'admin/data_management/migration_wizard.html', context)


@staff_member_required
def download_template(request):
    """
    Download template for selected module and year range.
    """
    module = request.GET.get('module')
    year_range = request.GET.get('years', '1')
    file_format = request.GET.get('format', 'xlsx')
    
    if not module or module not in _get_templates():
        messages.error(request, 'Invalid module selected')
        return redirect('nucleiq_admin:migration_wizard')
    
    try:
        if year_range == 'custom':
            year_range = '10'  # Default to 10 for custom
        
        wb = generate_template_with_years(module, year_range)
        
        # Generate filename
        module_config = _get_templates()[module]
        filename = f"{module}_import_template_{year_range}yr.xlsx"
        
        # Create response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        wb.save(response)
        return response
        
    except Exception as e:
        messages.error(request, f'Error generating template: {str(e)}')
        return redirect('nucleiq_admin:migration_wizard')


@staff_member_required
def upload_preview(request):
    """
    Handle file upload and preview data before import.
    """
    if request.method != 'POST':
        return redirect('nucleiq_admin:migration_wizard')
    
    tenant_id = request.POST.get('tenant')
    module = request.POST.get('module')
    year_range = request.POST.get('year_range', '1')
    uploaded_file = request.FILES.get('data_file')
    
    if not all([tenant_id, module, uploaded_file]):
        messages.error(request, 'Please fill in all required fields')
        return redirect('nucleiq_admin:migration_wizard')
    
    try:
        tenant = Tenant.objects.get(id=tenant_id)
    except Tenant.DoesNotExist:
        messages.error(request, 'Invalid tenant selected')
        return redirect('nucleiq_admin:migration_wizard')
    
    if module not in _get_templates():
        messages.error(request, 'Invalid module selected')
        return redirect('nucleiq_admin:migration_wizard')
    
    # Read and validate file
    try:
        import pandas as pd
        
        file_ext = uploaded_file.name.lower().split('.')[-1]
        
        if file_ext == 'csv':
            df = pd.read_csv(uploaded_file)
        elif file_ext in ['xlsx', 'xls']:
            df = pd.read_excel(uploaded_file)
        else:
            messages.error(request, 'Invalid file format. Please upload CSV or Excel file.')
            return redirect('nucleiq_admin:migration_wizard')
        
        # Get template for validation
        template = _get_templates()[module]['template']
        required_fields = [f.name for f in template.get_required_fields()]
        all_fields = [f.name for f in template.fields]
        
        # Check for required columns
        file_columns = [col.strip().lower() for col in df.columns]
        missing_required = []
        
        for req_field in required_fields:
            # Check with aliases too
            field_spec = next((f for f in template.fields if f.name == req_field), None)
            aliases = [req_field.lower()]
            if field_spec and hasattr(field_spec, 'aliases'):
                aliases.extend([a.lower() for a in field_spec.aliases])
            
            if not any(alias in file_columns for alias in aliases):
                missing_required.append(req_field)
        
        # Prepare preview data
        preview_rows = df.head(10).to_dict('records')
        total_rows = len(df)
        columns = list(df.columns)
        
        # Validation summary
        validation = {
            'total_rows': total_rows,
            'columns_found': len(columns),
            'missing_required': missing_required,
            'is_valid': len(missing_required) == 0,
        }
        
        # Store file in session for later import
        # Save to temp file
        temp_dir = tempfile.gettempdir()
        temp_path = os.path.join(temp_dir, f"nucleiq_import_{request.user.id}_{module}.{file_ext}")
        
        uploaded_file.seek(0)
        with open(temp_path, 'wb') as f:
            for chunk in uploaded_file.chunks():
                f.write(chunk)
        
        request.session['pending_import'] = {
            'tenant_id': tenant_id,
            'tenant_name': tenant.name,
            'module': module,
            'year_range': year_range,
            'file_path': temp_path,
            'file_name': uploaded_file.name,
            'total_rows': total_rows,
        }
        
        context = {
            'title': 'Preview Import Data',
            'tenant': tenant,
            'module': module,
            'module_name': _get_templates()[module]['template'].display_name,
            'file_name': uploaded_file.name,
            'preview_rows': preview_rows,
            'columns': columns,
            'validation': validation,
            'required_fields': required_fields,
            'all_fields': all_fields,
            'site_header': 'NucleiQ Platform Administration',
            'has_permission': True,
        }
        
        return render(request, 'admin/data_management/import_preview.html', context)
        
    except Exception as e:
        messages.error(request, f'Error reading file: {str(e)}')
        return redirect('nucleiq_admin:migration_wizard')


@staff_member_required
@require_http_methods(["POST"])
def execute_import(request):
    """
    Execute the actual import after preview confirmation.
    """
    pending_import = request.session.get('pending_import')
    
    if not pending_import:
        messages.error(request, 'No pending import found. Please start over.')
        return redirect('nucleiq_admin:migration_wizard')
    
    skip_duplicates = request.POST.get('skip_duplicates', 'true') == 'true'
    update_existing = request.POST.get('update_existing', 'false') == 'true'
    
    try:
        from .utils import process_import_file
        
        tenant_id = pending_import['tenant_id']
        module = pending_import['module']
        file_path = pending_import['file_path']
        
        tenant = Tenant.objects.get(id=tenant_id)
        
        # Create import job
        import_job = ImportJob.objects.create(
            tenant=tenant,
            module=module,
            status='IMPORTING',
            original_filename=pending_import['file_name'],
            file_type=file_path.split('.')[-1],
            total_rows=pending_import['total_rows'],
            skip_duplicates=skip_duplicates,
            update_existing=update_existing,
            created_by=request.user,
        )
        
        # Process the import
        with open(file_path, 'rb') as f:
            result = process_import_file(
                file_obj=f,
                module=module,
                tenant=tenant,
                skip_duplicates=skip_duplicates,
                update_existing=update_existing,
                job=import_job,
            )
        
        # Update job status
        import_job.status = 'COMPLETED' if result.get('success') else 'FAILED'
        import_job.successful_rows = result.get('success_count', 0)
        import_job.failed_rows = result.get('failed', 0)
        import_job.duplicate_rows = result.get('duplicates_skipped', 0)
        import_job.processed_rows = result.get('total', 0)
        import_job.validation_errors = result.get('errors', [])
        import_job.completed_at = timezone.now()
        import_job.save()
        
        # Clean up temp file
        if os.path.exists(file_path):
            os.remove(file_path)
        
        # Clear session
        del request.session['pending_import']
        
        if result.get('success'):
            messages.success(
                request, 
                f"Import completed! {result.get('success_count', 0)} records imported successfully."
            )
        else:
            messages.warning(
                request,
                f"Import completed with issues: {result.get('success_count', 0)} succeeded, "
                f"{result.get('failed', 0)} failed."
            )
        
        return redirect('nucleiq_admin:import_result', job_id=import_job.id)
        
    except Exception as e:
        messages.error(request, f'Import failed: {str(e)}')
        return redirect('nucleiq_admin:migration_wizard')


@staff_member_required
def import_result(request, job_id):
    """
    Show import result details.
    """
    try:
        job = ImportJob.objects.get(id=job_id)
    except ImportJob.DoesNotExist:
        messages.error(request, 'Import job not found')
        return redirect('nucleiq_admin:migration_wizard')
    
    context = {
        'title': 'Import Result',
        'job': job,
        'module_name': _get_templates().get(job.module, {}).get('template', {}).display_name if job.module in _get_templates() else job.module,
        'site_header': 'NucleiQ Platform Administration',
        'has_permission': True,
    }
    
    return render(request, 'admin/data_management/import_result.html', context)


@staff_member_required
def import_history(request):
    """
    Show import history with filtering.
    """
    tenant_id = request.GET.get('tenant')
    module = request.GET.get('module')
    status = request.GET.get('status')
    
    jobs = ImportJob.objects.all().order_by('-created_at')
    
    if tenant_id:
        jobs = jobs.filter(tenant_id=tenant_id)
    if module:
        jobs = jobs.filter(module=module)
    if status:
        jobs = jobs.filter(status=status)
    
    jobs = jobs[:50]  # Limit to last 50
    
    context = {
        'title': 'Import History',
        'jobs': jobs,
        'tenants': Tenant.objects.filter(is_active=True),
        'modules': [{'name': k, 'display_name': v['template'].display_name} for k, v in MODULE_TEMPLATES.items()],
        'selected_tenant': tenant_id,
        'selected_module': module,
        'selected_status': status,
        'site_header': 'NucleiQ Platform Administration',
        'has_permission': True,
    }
    
    return render(request, 'admin/data_management/import_history.html', context)


@staff_member_required
@require_http_methods(["POST"])
def rollback_import(request, job_id):
    """
    Rollback an import job.
    """
    try:
        job = ImportJob.objects.get(id=job_id)
        
        if job.status == 'ROLLED_BACK':
            messages.warning(request, 'This import has already been rolled back.')
            return redirect('nucleiq_admin:import_history')
        
        if not job.created_record_ids:
            messages.warning(request, 'No records to rollback.')
            return redirect('nucleiq_admin:import_history')
        
        # Perform rollback based on module
        from .utils import rollback_import_job
        
        result = rollback_import_job(job)
        
        if result.get('success'):
            job.status = 'ROLLED_BACK'
            job.save()
            messages.success(request, f"Rolled back {result.get('count', 0)} records.")
        else:
            messages.error(request, f"Rollback failed: {result.get('error', 'Unknown error')}")
        
    except ImportJob.DoesNotExist:
        messages.error(request, 'Import job not found')
    except Exception as e:
        messages.error(request, f'Rollback failed: {str(e)}')
    
    return redirect('nucleiq_admin:import_history')


@staff_member_required
def template_library(request):
    """
    Show all available templates for download.
    """
    templates = []
    
    for name, config in _get_templates().items():
        template = config['template']
        templates.append({
            'name': name,
            'display_name': template.display_name,
            'icon': config['icon'],
            'description': config['description'],
            'supports_historical': config.get('supports_historical', False),
            'required_fields': len(template.get_required_fields()),
            'optional_fields': len(template.get_optional_fields()),
            'instructions': template.instructions[:5] if template.instructions else [],
        })
    
    # Sort templates
    module_order = ['students', 'staff', 'classes', 'subjects', 'fee_structures', 
                    'fee_invoices', 'fee_payments', 'fee_allocations', 'fee_discounts',
                    'student_enrollments', 'attendance', 'transport', 'parents', 'user_accounts',
                    'student_photos', 'exam_results', 'exam_schedule', 'timetable',
                    'library_books', 'library_transactions', 'payroll_payments',
                    'hostel_allocations', 'inventory_items', 'certificates_issued',
                    'finance_journal_entries', 'helpdesk_tickets', 'lms_courses',
                    'lms_enrollments', 'idcards']
    templates.sort(key=lambda x: module_order.index(x['name']) if x['name'] in module_order else 999)
    
    context = {
        'title': 'Template Library',
        'templates': templates,
        'year_options': YEAR_RANGE_OPTIONS,
        'site_header': 'NucleiQ Platform Administration',
        'has_permission': True,
    }
    
    return render(request, 'admin/data_management/template_library.html', context)


@staff_member_required
def export_center(request):
    """
    Admin UI to export module data and perform full backups.
    """
    templates = []
    for name, config in _get_templates().items():
        template = config['template']
        templates.append({
            'name': name,
            'display_name': template.display_name,
            'icon': config['icon'],
            'description': config['description'],
            'supports_historical': config.get('supports_historical', False),
        })

    context = {
        'title': 'Export Center',
        'templates': templates,
        'site_header': 'NucleiQ Platform Administration',
        'has_permission': True,
    }

    return render(request, 'admin/data_management/export_center.html', context)
