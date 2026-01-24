"""
Enhanced Data Migration Views
Import, Export, Backup, and Migration functionality
Supports CSV, XLSX, and XLS file formats
"""

import io
import csv
import json
import zipfile
import re
import uuid
from datetime import datetime, date
from django.http import HttpResponse
from django.db import transaction, IntegrityError
from django.db.models import Max, Q
from django.utils import timezone
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from core.permissions import IsTenantAdmin

from .models import ImportJob, ImportFieldMapping
from .utils import (
    FileParser, DataTransformer, DataValidator,
    DateParser, PhoneValidator, EmailValidator,
    OPENPYXL_AVAILABLE, XLRD_AVAILABLE
)
from .templates import (
    get_template, get_all_templates, get_template_info,
    TEMPLATE_REGISTRY, FieldType
)

from students.models import Student, StudentEnrollment
from staff.models import Staff
from tenants.models import GradeLevel, Section, Subject, AcademicYear, Department
from fees.models import FeeCategory, FeeStructure, FeeAllocation
from attendance.models import AttendanceRecord
from transport.models import Route, Stop, StudentTransport
from exams.models import ExamTerm, Exam, ExamSchedule, ExamResult
from timetable.models import TimetableSlot
from library.models import Book, BookCopy, LibraryMember, BookIssue
from hostel.models import HostelBuilding, Room, Bed, HostelAllocation
from inventory.models import Item, ItemCategory
from certificates.models import CertificateTemplate, CertificateRequest, GeneratedCertificate
from finance.models import LedgerAccount, JournalEntry, JournalEntryLine, SalaryPayment
from helpdesk.models import HelpdeskTicket
from lms.models import LiveClass
from idcards.models import IDCardTemplate, IDCardRecord

import logging
logger = logging.getLogger(__name__)

# Try to import openpyxl for Excel support
try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    import openpyxl.utils
except ImportError:
    pass


class ModuleListView(APIView):
    """List all available import modules with their specifications"""
    # Allow any authenticated user to view available modules
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        """Get list of all importable modules"""
        modules = get_template_info()
        
        # Add format availability info
        formats_available = {
            'csv': True,
            'xlsx': OPENPYXL_AVAILABLE,
            'xls': XLRD_AVAILABLE,
        }
        
        return Response({
            'modules': modules,
            'supported_formats': formats_available,
            'max_file_size_mb': 10,
        })


class ModuleFieldsView(APIView):
    """Get detailed field information for a specific module"""
    # Allow any authenticated user to view field information
    permission_classes = [IsAuthenticated]
    
    def get(self, request, module):
        """Get field specifications for a module"""
        template = get_template(module)
        
        if not template:
            return Response(
                {'error': f'Unknown module: {module}'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        return Response(template.to_dict())


class DownloadTemplateView(APIView):
    """Download Excel/CSV template for a specific module"""
    # Allow any authenticated user to download templates
    permission_classes = [IsAuthenticated]
    
    def get(self, request, module):
        format_type = request.query_params.get('format', 'xlsx')
        
        template = get_template(module)
        if not template:
            return Response(
                {'error': f'Invalid module: {module}'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        if format_type == 'xlsx':
            if not OPENPYXL_AVAILABLE:
                return Response(
                    {'error': 'Excel support not available. Please download CSV.'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            return self._generate_excel_template(module, template)
        else:
            return self._generate_csv_template(module, template)
    
    def _generate_excel_template(self, module, template):
        """Generate Excel template with formatting and instructions"""
        wb = Workbook()
        ws = wb.active
        ws.title = template.display_name
        
        # Header styling
        required_font = Font(bold=True, color='FFFFFF')
        optional_font = Font(bold=True, color='FFFFFF')
        required_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')  # Blue
        optional_fill = PatternFill(start_color='6B7280', end_color='6B7280', fill_type='solid')  # Gray
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # Write header row using internal field names (required by parser)
        all_fields = template.fields
        for col, field_spec in enumerate(all_fields, 1):
            cell = ws.cell(row=1, column=col, value=field_spec.name)
            # Visual hint: color the header to indicate required/optional
            cell.font = required_font if field_spec.required else optional_font
            cell.fill = required_fill if field_spec.required else optional_fill
            cell.border = thin_border
            cell.alignment = center_alignment

            # Set column width based on display name for readability
            col_letter = openpyxl.utils.get_column_letter(col)
            ws.column_dimensions[col_letter].width = max(15, len(field_spec.display_name) + 2)

        # Write sample data row (row 2)
        for col, field_spec in enumerate(all_fields, 1):
            cell = ws.cell(row=2, column=col, value=field_spec.sample_value)
            cell.border = thin_border
        
        # Create Instructions sheet
        instructions_ws = wb.create_sheet('Instructions')
        
        row = 1
        for instruction in template.instructions:
            instructions_ws.cell(row=row, column=1, value=instruction)
            if instruction.startswith('📋') or instruction.startswith('💡') or instruction.startswith('⚠️') or instruction.startswith('✅') or instruction.startswith('⚪'):
                instructions_ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

        required_fields = [f.display_name for f in all_fields if f.required]
        optional_fields = [f.display_name for f in all_fields if not f.required]

        row += 1
        instructions_ws.cell(row=row, column=1, value='✅ REQUIRED FIELDS (must be filled):')
        instructions_ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1
        instructions_ws.cell(row=row, column=1, value=', '.join(required_fields) if required_fields else 'None')
        row += 1
        instructions_ws.cell(row=row, column=1, value='⚪ OPTIONAL FIELDS (can be left blank):')
        instructions_ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1
        instructions_ws.cell(row=row, column=1, value=', '.join(optional_fields) if optional_fields else 'None')

        row += 2
        instructions_ws.cell(row=row, column=1, value='📝 FIELD SPECIFICATIONS:')
        instructions_ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 2
        
        # Field specification table
        headers = ['Field Name', 'Type', 'Required', 'Max Length', 'Description', 'Sample']
        for col, header in enumerate(headers, 1):
            cell = instructions_ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='E5E7EB', fill_type='solid')
        row += 1
        
        for field_spec in all_fields:
            instructions_ws.cell(row=row, column=1, value=field_spec.name)
            instructions_ws.cell(row=row, column=2, value=field_spec.field_type.value)
            instructions_ws.cell(row=row, column=3, value='Yes' if field_spec.required else 'No')
            instructions_ws.cell(row=row, column=4, value=field_spec.max_length or '-')
            instructions_ws.cell(row=row, column=5, value=field_spec.description)
            instructions_ws.cell(row=row, column=6, value=field_spec.sample_value)
            row += 1
        
        # Adjust column widths
        instructions_ws.column_dimensions['A'].width = 25
        instructions_ws.column_dimensions['B'].width = 12
        instructions_ws.column_dimensions['C'].width = 10
        instructions_ws.column_dimensions['D'].width = 12
        instructions_ws.column_dimensions['E'].width = 50
        instructions_ws.column_dimensions['F'].width = 20
        
        # Return response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{module}_import_template.xlsx"'
        wb.save(response)
        return response
    
    def _generate_csv_template(self, module, template):
        """Generate CSV template"""
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{module}_import_template.csv"'
        
        # Add BOM for Excel compatibility
        response.write('\ufeff')
        
        writer = csv.writer(response)
        
        # Write internal field names as header (required by parser)
        field_names = [f.name for f in template.fields]
        writer.writerow(field_names)

        # Write sample data (second row)
        sample_data = [f.sample_value for f in template.fields]
        writer.writerow(sample_data)
        
        return response


class ValidateDataView(APIView):
    """Validate uploaded file without importing - Admin only"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        """Validate file and return preview with any errors"""
        module = request.data.get('module')
        file = request.FILES.get('file')
        
        template = get_template(module)
        if not template:
            return Response(
                {'error': f'Invalid module: {module}'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        if not file:
            return Response(
                {'error': 'No file provided'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        try:
            # Parse file
            data, _, file_type = FileParser.parse(file, file.name)
            
            if not data:
                return Response({
                    'valid': False,
                    'error': 'No data found in file',
                    'total_rows': 0
                })
            
            # Validate data
            validation_result = self._validate_data(module, template, data, request.user.tenant)
            
            return Response(validation_result)
            
        except Exception as e:
            logger.exception(f"Validation error: {e}")
            return Response(
                {'error': str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
    
    def _validate_data(self, module, template, data, tenant):
        """Validate all rows in the data"""
        errors = []
        warnings = []
        duplicates = []
        preview_rows = []
        
        # Get unique field for duplicate detection
        unique_field = template.unique_field
        
        # Get transformer based on module
        transformer = self._get_transformer(module)
        
        for idx, row in enumerate(data, start=2):
            row_num = row.get('_row_number', idx)
            
            # Transform the row
            transformed = transformer(row) if transformer else row
            
            # Validate
            row_errors = self._validate_row(module, transformed, row_num)
            errors.extend(row_errors)
            
            # Check for duplicates
            if unique_field and transformed.get(unique_field):
                existing = self._check_duplicate(module, unique_field, transformed[unique_field], tenant)
                if existing:
                    duplicates.append({
                        'row': row_num,
                        'field': unique_field,
                        'value': transformed[unique_field],
                        'existing_id': existing.id if hasattr(existing, 'id') else None
                    })
            
            # Add to preview (first 10 rows)
            if idx <= 11:  # First 10 data rows
                preview_rows.append({
                    'row_number': row_num,
                    'data': {k: str(v) if v else '' for k, v in transformed.items() if not k.startswith('_')},
                    'has_errors': len(row_errors) > 0
                })
        
        return {
            'valid': len(errors) == 0,
            'total_rows': len(data),
            'error_count': len(errors),
            'duplicate_count': len(duplicates),
            'warning_count': len(warnings),
            'errors': errors[:50],  # Limit errors returned
            'warnings': warnings[:20],
            'duplicates': duplicates[:20],
            'preview': preview_rows,
        }
    
    def _get_transformer(self, module):
        """Get the appropriate data transformer for a module"""
        transformers = {
            'students': DataTransformer.transform_student,
            'staff': DataTransformer.transform_staff,
            'classes': DataTransformer.transform_class_section,
        }
        return transformers.get(module)
    
    def _validate_row(self, module, data, row_num):
        """Validate a single row"""
        validators = {
            'students': DataValidator.validate_student,
            'staff': DataValidator.validate_staff,
            'classes': DataValidator.validate_class_section,
        }
        
        validator = validators.get(module)
        if validator:
            return validator(data, row_num)
        return []
    
    def _check_duplicate(self, module, field, value, tenant):
        """Check if a record with the same unique field exists"""
        models = {
            'students': Student,
            'staff': Staff,
            'subjects': Subject,
        }
        
        model = models.get(module)
        if model:
            try:
                filter_kwargs = {field: value, 'tenant': tenant, 'is_deleted': False}
                return model.objects.filter(**filter_kwargs).first()
            except:
                return None
        return None


class ImportDataView(APIView):
    """Import data from CSV/XLSX/XLS file - Admin only"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        """Import data from uploaded file"""
        module = request.data.get('module')
        file = request.FILES.get('file')
        skip_duplicates = request.data.get('skip_duplicates', 'true') == 'true'
        update_existing = request.data.get('update_existing', 'false') == 'true'

        # Store request user for downstream helpers
        self._request_user = request.user
        
        template = get_template(module)
        if not template:
            return Response(
                {'error': f'Invalid module: {module}'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        if not file:
            return Response(
                {'error': 'No file provided'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Special handling for student photo ZIP uploads
        if module == 'student_photos' and file.name.lower().endswith('.zip'):
            return Response(self._process_photo_zip(file, request.user.tenant))
        
        tenant = request.user.tenant
        
        # Create import job for tracking
        import_job = ImportJob.objects.create(
            tenant=tenant,
            module=module,
            status='VALIDATING',
            original_filename=file.name,
            file_type=file.name.split('.')[-1].lower(),
            skip_duplicates=skip_duplicates,
            update_existing=update_existing,
            created_by=request.user
        )
        
        try:
            # Parse file
            data, _, file_type = FileParser.parse(file, file.name)
            import_job.total_rows = len(data)
            import_job.save()
            
            if not data:
                import_job.status = 'FAILED'
                import_job.validation_errors = ['No data found in file']
                import_job.save()
                return Response({
                    'success': False,
                    'error': 'No data found in file',
                    'job_id': str(import_job.id)
                })
            
            # Process import
            import_job.status = 'IMPORTING'
            import_job.save()
            
            result = self._process_import(
                module, template, data, tenant, 
                skip_duplicates, update_existing, import_job
            )
            
            # Update job status
            import_job.status = 'COMPLETED' if result['failed'] == 0 else 'COMPLETED'
            import_job.successful_rows = result['success']
            import_job.failed_rows = result['failed']
            import_job.duplicate_rows = result['duplicates_skipped']
            import_job.validation_errors = result['errors'][:100]
            import_job.completed_at = timezone.now()
            import_job.save()
            
            return Response({
                'success': True,
                'job_id': str(import_job.id),
                **result
            })
            
        except Exception as e:
            logger.exception(f"Import error: {e}")
            import_job.status = 'FAILED'
            import_job.validation_errors = [str(e)]
            import_job.save()
            
            return Response(
                {'error': str(e), 'job_id': str(import_job.id)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
    
    def _process_import(self, module, template, data, tenant, skip_duplicates, update_existing, import_job):
        """Process the import with transaction management"""
        success = 0
        failed = 0
        duplicates_skipped = 0
        updated = 0
        errors = []
        created_ids = []
        
        transformer = self._get_transformer(module)
        unique_field = template.unique_field
        
        for idx, row in enumerate(data):
            row_num = row.get('_row_number', idx + 2)

            try:
                # Use a nested savepoint so a failure in one row doesn't abort the whole import
                with transaction.atomic():
                    # Transform the row
                    transformed = transformer(row) if transformer else row

                    # Check for duplicates
                    existing = None
                    if unique_field and transformed.get(unique_field):
                        existing = self._find_existing(module, unique_field, transformed[unique_field], tenant)

                    if existing:
                        if update_existing:
                            self._update_record(module, existing, transformed, tenant)
                            updated += 1
                        elif skip_duplicates:
                            duplicates_skipped += 1
                            # commit savepoint and continue
                            continue
                        else:
                            errors.append(f'Row {row_num}: Duplicate {unique_field}: {transformed[unique_field]}')
                            failed += 1
                            continue
                    else:
                        # Create new record
                        record = self._create_record(module, transformed, tenant)
                        if record:
                            created_ids.append(str(record.id))
                            success += 1
                        else:
                            failed += 1
                            errors.append(f'Row {row_num}: Failed to create record')

                    # Update progress
                    import_job.processed_rows = idx + 1
                    if (idx + 1) % 50 == 0:  # Save every 50 rows
                        import_job.save()

            except IntegrityError as e:
                # Roll back to savepoint implicitly and record the error for this row
                failed += 1
                errors.append(f'Row {row_num}: {str(e)}')
                logger.warning(f"Error importing row {row_num}: {e}")
                continue
            except Exception as e:
                failed += 1
                errors.append(f'Row {row_num}: {str(e)}')
                logger.warning(f"Error importing row {row_num}: {e}")
        
        # Store created IDs for potential rollback
        import_job.created_record_ids = created_ids
        
        return {
            'success': success,
            'failed': failed,
            'updated': updated,
            'duplicates_skipped': duplicates_skipped,
            'total': len(data),
            'errors': errors[:50]
        }

    def _normalize_grade_name(self, value):
        if value is None:
            return ''
        raw = re.sub(r'\s+', ' ', str(value)).strip()
        if not raw:
            return ''

        lower = raw.lower()
        lower = re.sub(r'^(class|grade|std|standard|year)\s+', '', lower)
        lower = lower.replace('-', ' ')
        lower = re.sub(r'\s+', ' ', lower).strip()

        roman_map = {
            'i': '1', 'ii': '2', 'iii': '3', 'iv': '4', 'v': '5',
            'vi': '6', 'vii': '7', 'viii': '8', 'ix': '9', 'x': '10',
            'xi': '11', 'xii': '12'
        }
        if lower in roman_map:
            lower = roman_map[lower]

        return lower

    def _canonical_grade_fields(self, value):
        raw = re.sub(r'\s+', ' ', str(value or '')).strip()
        if not raw:
            return '', ''

        normalized = self._normalize_grade_name(raw)
        if normalized.isdigit():
            return f"Class {normalized}", normalized

        if normalized in {'kg', 'lkg', 'ukg'}:
            name = normalized.upper()
            return name, name

        short_name = re.sub(r'^(Class|Grade|Std|Standard)\s+', '', raw, flags=re.IGNORECASE)
        short_name = short_name.strip() if short_name else raw
        return raw, short_name

    def _find_grade_level(self, tenant, class_name):
        if not class_name:
            return None

        raw = re.sub(r'\s+', ' ', str(class_name)).strip()
        normalized = self._normalize_grade_name(raw)

        q = Q(name__iexact=raw)
        if normalized:
            q |= Q(name__iexact=normalized) | Q(short_name__iexact=normalized)
            if normalized.isdigit():
                q |= Q(name__iexact=f"Class {normalized}") | Q(name__iexact=f"Grade {normalized}") | Q(name__iexact=f"Std {normalized}")
            if normalized in {'kg', 'lkg', 'ukg'}:
                q |= Q(name__iexact=normalized.upper()) | Q(short_name__iexact=normalized.upper())

        return GradeLevel.objects.filter(
            tenant=tenant,
            is_deleted=False
        ).filter(q).first()

    def _normalize_section_name(self, value):
        if value is None:
            return ''
        raw = re.sub(r'\s+', ' ', str(value)).strip()
        if not raw:
            return ''

        lower = raw.lower()
        lower = re.sub(r'^(section|sec)\s+', '', lower)
        lower = lower.replace('-', ' ')
        lower = re.sub(r'\s+', ' ', lower).strip()

        if len(lower) == 1:
            return lower.upper()

        return lower.title()

    def _find_section(self, tenant, grade_level, section_name):
        if not section_name or not grade_level:
            return None

        raw = re.sub(r'\s+', ' ', str(section_name)).strip()
        normalized = self._normalize_section_name(raw)

        q = Q(name__iexact=raw)
        if normalized:
            q |= Q(name__iexact=normalized)

        return Section.objects.filter(
            tenant=tenant,
            grade_level=grade_level,
            is_deleted=False
        ).filter(q).first()

    def _get_academic_year(self, tenant, year_name=None):
        if year_name:
            year = AcademicYear.objects.filter(
                tenant=tenant,
                name__iexact=str(year_name).strip(),
                is_deleted=False
            ).first()
            if year:
                return year

        year = AcademicYear.objects.filter(
            tenant=tenant,
            is_active=True,
            is_deleted=False
        ).first()

        if not year:
            year = AcademicYear.objects.filter(
                tenant=tenant,
                is_deleted=False
            ).order_by('-start_date').first()

        return year

    def _parse_time(self, value, default=None):
        if value is None or value == '':
            return default
        if hasattr(value, 'strftime') and hasattr(value, 'hour'):
            return value
        value = str(value).strip()
        for fmt in ['%H:%M', '%H:%M:%S']:
            try:
                return datetime.strptime(value, fmt).time()
            except ValueError:
                continue
        return default

    def _parse_datetime(self, value, default=None):
        if value is None or value == '':
            return default
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        value = str(value).strip()
        for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M']:
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                continue
        parsed_date = DateParser.parse(value)
        if parsed_date:
            return datetime.combine(parsed_date, datetime.min.time())
        return default

    def _process_photo_zip(self, file, tenant):
        """Process a ZIP file containing student photos named by admission number"""
        success = 0
        failed = 0
        errors = []

        try:
            with zipfile.ZipFile(file) as z:
                for filename in z.namelist():
                    if filename.endswith('/') or not filename.lower().endswith(('.jpg', '.jpeg', '.png')):
                        continue

                    name_part = filename.split('/')[-1]
                    admission_number = name_part.rsplit('.', 1)[0]

                    try:
                        student = Student.objects.filter(
                            tenant=tenant,
                            admission_number=admission_number,
                            is_deleted=False
                        ).first()

                        if student:
                            from django.core.files.base import ContentFile
                            photo_data = z.read(filename)
                            ext = name_part.rsplit('.', 1)[1]
                            student.photo.save(
                                f"{admission_number}.{ext}",
                                ContentFile(photo_data),
                                save=True
                            )
                            success += 1
                        else:
                            failed += 1
                            errors.append(f"Student {admission_number} not found (File: {filename})")
                    except Exception as e:
                        failed += 1
                        errors.append(f"Error processing {filename}: {str(e)}")

        except Exception as e:
            return {'success': 0, 'failed': 0, 'duplicates': [], 'errors': [f"Invalid ZIP: {str(e)}"]}

        return {
            'success': success,
            'failed': failed,
            'duplicates': [],
            'errors': errors[:20]
        }

    def _create_fee_allocation(self, data, tenant):
        admission_number = data.get('admission_number')
        fee_type = data.get('fee_type')
        amount = data.get('amount')
        academic_year_name = data.get('academic_year')

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        fee_category = FeeCategory.objects.filter(
            tenant=tenant,
            name__iexact=str(fee_type).strip()
        ).first()
        if not fee_category:
            raise ValueError(f"Fee Category '{fee_type}' not found")

        enrollment = student.get_current_enrollment() if hasattr(student, 'get_current_enrollment') else None
        class_name = data.get('class_name') or (enrollment.class_assigned.name if enrollment and enrollment.class_assigned else None)

        if not class_name:
            raise ValueError(f"Class not found for student '{admission_number}'")

        academic_year = self._get_academic_year(tenant, academic_year_name)
        if not academic_year:
            raise ValueError("No academic year found")

        fee_structure = FeeStructure.objects.filter(
            tenant=tenant,
            academic_year=academic_year,
            class_level=class_name,
            category=fee_category
        ).first()

        if not fee_structure:
            raise ValueError(f"Fee Structure not found for {class_name} - {fee_type} ({academic_year.name})")

        custom_amount = None
        if amount is not None and str(amount).strip() != '':
            try:
                amount_value = float(amount)
                if amount_value != float(fee_structure.amount):
                    custom_amount = amount_value
            except Exception:
                pass

        allocation, _ = FeeAllocation.objects.get_or_create(
            tenant=tenant,
            student=student,
            fee_structure=fee_structure,
            defaults={
                'custom_amount': custom_amount
            }
        )

        if custom_amount is not None:
            allocation.custom_amount = custom_amount
            allocation.save(update_fields=['custom_amount'])

        return allocation

    def _create_student_photo(self, data, tenant):
        admission_number = data.get('admission_number')
        photo_filename = data.get('photo_filename')

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        if photo_filename:
            raise ValueError("Photo filename provided without ZIP upload. Please upload a ZIP of photos.")

        return student

    def _create_exam_result(self, data, tenant):
        admission_number = data.get('admission_number')
        exam_name = data.get('exam_name')
        subject_code = data.get('subject_code')
        marks_obtained = data.get('marks_obtained')
        max_marks = data.get('max_marks')
        academic_year_name = data.get('academic_year')
        exam_date = DateParser.parse(data.get('exam_date'))

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        subject = Subject.objects.filter(
            tenant=tenant,
            code__iexact=str(subject_code).strip()
        ).first()
        if not subject:
            raise ValueError(f"Subject '{subject_code}' not found")

        enrollment = student.get_current_enrollment() if hasattr(student, 'get_current_enrollment') else None
        grade_level = enrollment.section.grade_level if enrollment and enrollment.section else None
        section = enrollment.section if enrollment else None

        if not grade_level:
            raise ValueError(f"Student '{admission_number}' has no active class/section")

        academic_year = self._get_academic_year(tenant, academic_year_name)
        if not academic_year:
            raise ValueError("No academic year found")

        term_name = f"{exam_name} {academic_year.name}" if exam_name else f"Exam {academic_year.name}"
        term = ExamTerm.objects.filter(
            tenant=tenant,
            name=term_name,
            academic_year=academic_year,
            is_deleted=False
        ).first()
        if not term:
            term = ExamTerm.objects.create(
                tenant=tenant,
                name=term_name,
                term_type='OTHER',
                academic_year=academic_year,
                start_date=exam_date or academic_year.start_date,
                end_date=exam_date or academic_year.end_date,
                is_active=True
            )

        exam = Exam.objects.filter(
            tenant=tenant,
            exam_term=term,
            subject=subject,
            grade_level=grade_level,
            is_deleted=False
        ).first()
        if not exam:
            total_marks = float(max_marks) if max_marks else 100
            exam = Exam.objects.create(
                tenant=tenant,
                name=exam_name or f"{subject.name} Exam",
                exam_term=term,
                subject=subject,
                grade_level=grade_level,
                total_marks=total_marks,
                passing_marks=total_marks * 0.4,
                duration_minutes=60,
                status='COMPLETED'
            )
            if section:
                exam.sections.add(section)

        if section and not exam.sections.filter(id=section.id).exists():
            exam.sections.add(section)

        result, _ = ExamResult.objects.get_or_create(
            tenant=tenant,
            exam=exam,
            student=student,
            section=section or exam.sections.first(),
            defaults={
                'marks_obtained': float(marks_obtained or 0),
                'status': 'DRAFT'
            }
        )

        if marks_obtained is not None:
            result.marks_obtained = float(marks_obtained)
            result.save(update_fields=['marks_obtained'])

        return result

    def _create_exam_schedule(self, data, tenant):
        exam_name = data.get('exam_name')
        class_name = data.get('class_name')
        subject_code = data.get('subject_code')
        exam_date = DateParser.parse(data.get('exam_date'))
        start_time = self._parse_time(data.get('start_time')) or self._parse_time('09:00')
        end_time = self._parse_time(data.get('end_time')) or self._parse_time('12:00')
        max_marks = data.get('max_marks')
        section_name = data.get('section_name')

        grade_level = self._find_grade_level(tenant, class_name)
        if not grade_level:
            raise ValueError(f"Class '{class_name}' not found")

        section = self._find_section(tenant, grade_level, section_name) or Section.objects.filter(
            tenant=tenant,
            grade_level=grade_level,
            is_deleted=False
        ).first()
        if not section:
            raise ValueError(f"Section not found for class '{class_name}'")

        subject = Subject.objects.filter(
            tenant=tenant,
            code__iexact=str(subject_code).strip()
        ).first()
        if not subject:
            raise ValueError(f"Subject '{subject_code}' not found")

        academic_year = self._get_academic_year(tenant, data.get('academic_year'))
        if not academic_year:
            raise ValueError("No academic year found")

        term_name = f"{exam_name} {academic_year.name}" if exam_name else f"Exam {academic_year.name}"
        term = ExamTerm.objects.filter(
            tenant=tenant,
            name=term_name,
            academic_year=academic_year,
            is_deleted=False
        ).first()
        if not term:
            term = ExamTerm.objects.create(
                tenant=tenant,
                name=term_name,
                term_type='OTHER',
                academic_year=academic_year,
                start_date=exam_date or academic_year.start_date,
                end_date=exam_date or academic_year.end_date,
                is_active=True
            )

        exam = Exam.objects.filter(
            tenant=tenant,
            exam_term=term,
            subject=subject,
            grade_level=grade_level,
            is_deleted=False
        ).first()
        if not exam:
            total_marks = float(max_marks) if max_marks else 100
            exam = Exam.objects.create(
                tenant=tenant,
                name=exam_name or f"{subject.name} Exam",
                exam_term=term,
                subject=subject,
                grade_level=grade_level,
                total_marks=total_marks,
                passing_marks=total_marks * 0.4,
                duration_minutes=60,
                status='SCHEDULED'
            )

        if not exam.sections.filter(id=section.id).exists():
            exam.sections.add(section)

        schedule, _ = ExamSchedule.objects.get_or_create(
            tenant=tenant,
            exam=exam,
            section=section,
            defaults={
                'exam_date': exam_date or date.today(),
                'start_time': start_time,
                'end_time': end_time,
                'room': data.get('room') or ''
            }
        )

        return schedule

    def _create_timetable_slot(self, data, tenant):
        class_name = data.get('class_name')
        section_name = data.get('section_name')
        subject_code = data.get('subject_code')
        day_of_week = data.get('day_of_week')
        start_time = self._parse_time(data.get('start_time'))
        end_time = self._parse_time(data.get('end_time'))
        period_number = data.get('period_number')
        teacher_employee_id = data.get('teacher_employee_id')
        room_number = data.get('room_number')

        grade_level = self._find_grade_level(tenant, class_name)
        if not grade_level:
            raise ValueError(f"Class '{class_name}' not found")

        section = self._find_section(tenant, grade_level, section_name)
        if not section:
            raise ValueError(f"Section '{section_name}' not found for class '{class_name}'")

        subject = Subject.objects.filter(
            tenant=tenant,
            code__iexact=str(subject_code).strip()
        ).first()
        if not subject:
            raise ValueError(f"Subject '{subject_code}' not found")

        academic_year = self._get_academic_year(tenant, data.get('academic_year'))
        if not academic_year:
            raise ValueError("No academic year found")

        teacher = None
        if teacher_employee_id:
            teacher = Staff.objects.filter(
                tenant=tenant,
                employee_id=teacher_employee_id,
                is_deleted=False
            ).first()

        day_map = {
            'MON': 'MONDAY', 'TUE': 'TUESDAY', 'WED': 'WEDNESDAY',
            'THU': 'THURSDAY', 'FRI': 'FRIDAY', 'SAT': 'SATURDAY', 'SUN': 'SUNDAY'
        }
        normalized_day = day_map.get(str(day_of_week).upper(), str(day_of_week).upper())

        if not start_time or not end_time:
            raise ValueError("Start time and end time are required")

        slot = TimetableSlot.objects.create(
            tenant=tenant,
            academic_year=academic_year,
            section=section,
            subject=subject,
            teacher=teacher,
            day_of_week=normalized_day,
            start_time=start_time,
            end_time=end_time,
            room=room_number or '',
            period_number=int(period_number) if str(period_number).isdigit() else None
        )

        return slot

    def _create_library_book(self, data, tenant):
        title = data.get('title')
        author = data.get('author')
        isbn = data.get('isbn')
        publisher = data.get('publisher')
        category = data.get('category')
        copies_total = data.get('copies_total')
        location = data.get('location')
        published_year = data.get('published_year')

        total_copies = int(float(copies_total)) if copies_total not in [None, ''] else 0
        year_published = int(float(published_year)) if published_year not in [None, ''] else None

        book = None
        if isbn:
            book = Book.objects.filter(
                tenant=tenant,
                isbn=str(isbn).strip(),
                is_deleted=False
            ).first()
        if not book:
            book = Book.objects.filter(
                tenant=tenant,
                title=title,
                author=author,
                is_deleted=False
            ).first()

        if not book:
            book = Book.objects.create(
                tenant=tenant,
                isbn=str(isbn).strip() if isbn else '',
                title=title,
                author=author or '',
                publisher=publisher or '',
                category=category or 'General',
                shelf_location=location or '',
                year_published=year_published,
                total_copies=total_copies,
                available_copies=total_copies,
            )

        if book:
            book.author = author or book.author
            book.publisher = publisher or book.publisher
            book.category = category or book.category
            book.shelf_location = location or book.shelf_location
            if year_published:
                book.year_published = year_published
            if total_copies:
                book.total_copies = total_copies
                book.available_copies = total_copies
            book.save()

        return book

    def _create_library_transaction(self, data, tenant):
        transaction_id = data.get('transaction_id')
        admission_number = data.get('admission_number')
        isbn = data.get('isbn')
        issue_date = DateParser.parse(data.get('issue_date'))
        due_date = DateParser.parse(data.get('due_date'))
        return_date = DateParser.parse(data.get('return_date'))
        fine_amount = data.get('fine_amount')
        status_value = data.get('status')

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        member, _ = LibraryMember.objects.get_or_create(
            tenant=tenant,
            student=student,
            defaults={'member_type': 'STUDENT'}
        )

        book = Book.objects.filter(
            tenant=tenant,
            isbn=str(isbn).strip()
        ).first()
        if not book:
            raise ValueError(f"Book with ISBN '{isbn}' not found")

        barcode = str(transaction_id or f"{isbn}-{student.admission_number}")
        copy, _ = BookCopy.objects.get_or_create(
            tenant=tenant,
            book=book,
            barcode=barcode,
            defaults={'status': 'AVAILABLE'}
        )

        issue = BookIssue.objects.create(
            tenant=tenant,
            copy=copy,
            member=member,
            issued_date=self._parse_datetime(issue_date) or timezone.now(),
            due_date=self._parse_datetime(due_date) or timezone.now(),
            returned_date=self._parse_datetime(return_date) if return_date else None,
            fine_amount=float(fine_amount or 0),
            status=status_value or 'ISSUED'
        )

        return issue

    def _create_payroll_payment(self, data, tenant):
        payment_reference = data.get('payment_reference')
        employee_id = data.get('employee_id')
        period_start = DateParser.parse(data.get('period_start'))
        payment_date = DateParser.parse(data.get('payment_date'))
        gross_amount = float(data.get('gross_amount') or 0)
        deductions = float(data.get('deductions') or 0)
        net_amount = float(data.get('net_amount') or 0)

        staff = Staff.objects.filter(
            tenant=tenant,
            employee_id=employee_id,
            is_deleted=False
        ).first()
        if not staff:
            raise ValueError(f"Staff '{employee_id}' not found")

        month = period_start or payment_date or date.today()

        payment, _ = SalaryPayment.objects.get_or_create(
            tenant=tenant,
            payment_number=payment_reference,
            defaults={
                'payment_date': payment_date or date.today(),
                'staff': staff,
                'month': date(month.year, month.month, 1),
                'basic_salary': gross_amount,
                'allowances': 0,
                'deductions': deductions,
                'net_salary': net_amount,
                'status': 'PAID' if payment_date else 'PENDING',
                'remarks': data.get('remarks') or ''
            }
        )

        return payment

    def _create_hostel_allocation(self, data, tenant):
        admission_number = data.get('admission_number')
        hostel_name = data.get('hostel_name')
        room_number = data.get('room_number')
        bed_number = data.get('bed_number')
        start_date = DateParser.parse(data.get('start_date')) or date.today()
        end_date = DateParser.parse(data.get('end_date'))
        status_value = data.get('status')

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        building, _ = HostelBuilding.objects.get_or_create(
            tenant=tenant,
            name=hostel_name,
            defaults={'building_type': 'BOYS', 'is_active': True}
        )

        room, _ = Room.objects.get_or_create(
            tenant=tenant,
            building=building,
            room_number=room_number,
            defaults={'capacity': 4}
        )

        bed, _ = Bed.objects.get_or_create(
            tenant=tenant,
            room=room,
            bed_number=bed_number or 'A'
        )

        allocation, _ = HostelAllocation.objects.get_or_create(
            tenant=tenant,
            student=student,
            bed=bed,
            defaults={
                'start_date': start_date,
                'end_date': end_date,
                'is_active': str(status_value).upper() != 'INACTIVE'
            }
        )

        return allocation

    def _create_inventory_item(self, data, tenant):
        item_code = data.get('item_code')
        item_name = data.get('item_name')
        category_name = data.get('category')
        quantity = data.get('quantity')
        unit_cost = data.get('unit_cost')
        purchase_date = data.get('purchase_date')
        vendor_name = data.get('vendor')
        location = data.get('location')
        condition = data.get('condition')

        category = None
        if category_name:
            category, _ = ItemCategory.objects.get_or_create(
                tenant=tenant,
                name=category_name
            )

        item, _ = Item.objects.get_or_create(
            tenant=tenant,
            sku=str(item_code).strip() if item_code else '',
            defaults={
                'name': item_name,
                'category': category,
                'current_stock': int(float(quantity or 0)),
                'cost_price': float(unit_cost or 0),
                'price': float(unit_cost or 0),
                'description': f"Location: {location or ''} | Condition: {condition or ''}".strip(),
                'is_active': True
            }
        )

        return item

    def _create_certificate_issued(self, data, tenant):
        certificate_number = data.get('certificate_number')
        admission_number = data.get('admission_number')
        certificate_type = data.get('certificate_type')
        issue_date = DateParser.parse(data.get('issue_date')) or date.today()

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        template, _ = CertificateTemplate.objects.get_or_create(
            tenant=tenant,
            name=certificate_type,
            defaults={'content': certificate_type or 'Certificate'}
        )

        request = CertificateRequest.objects.create(
            tenant=tenant,
            student=student,
            template=template,
            reason=data.get('remarks') or '',
            status='GENERATED'
        )

        generated = GeneratedCertificate.objects.create(
            tenant=tenant,
            request=request,
            certificate_number=certificate_number,
            issued_date=issue_date,
            content_snapshot=template.content,
            verified=True
        )

        return generated

    def _create_finance_journal_entry(self, data, tenant):
        entry_number = data.get('entry_number')
        entry_date = DateParser.parse(data.get('entry_date')) or date.today()
        account_code = data.get('account_code')
        account_name = data.get('account_name')
        debit = float(data.get('debit') or 0)
        credit = float(data.get('credit') or 0)
        narration = data.get('narration') or data.get('description') or ''
        reference = data.get('reference')

        if not entry_number:
            raise ValueError("Entry number is required")

        account = LedgerAccount.objects.filter(
            tenant=tenant,
            code=str(account_code).strip()
        ).first()

        if not account:
            account_type = 'INCOME' if credit > 0 else 'EXPENSE'
            account = LedgerAccount.objects.create(
                tenant=tenant,
                code=str(account_code).strip(),
                name=account_name or f"Account {account_code}",
                account_type=account_type,
                is_active=True
            )

        entry, _ = JournalEntry.objects.get_or_create(
            tenant=tenant,
            entry_number=entry_number,
            defaults={
                'entry_date': entry_date,
                'description': narration or f"Journal Entry {entry_number}",
                'reference_type': 'IMPORT',
                'reference_id': None,
                'status': 'DRAFT',
                'is_posted': False,
                'created_by': self._request_user
            }
        )

        line = JournalEntryLine.objects.create(
            entry=entry,
            account=account,
            description=narration or account.name,
            debit_amount=debit,
            credit_amount=credit
        )

        return line

    def _apply_fee_discount(self, data, tenant):
        admission_number = data.get('admission_number')
        discount_type = data.get('discount_type')
        value = float(data.get('value') or 0)
        reason = data.get('reason')

        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        if not student:
            raise ValueError(f"Student '{admission_number}' not found")

        allocation = FeeAllocation.objects.filter(
            tenant=tenant,
            student=student
        ).first()
        if not allocation:
            raise ValueError(f"No fee allocation found for student '{admission_number}'")

        if str(discount_type).upper() == 'PERCENT':
            allocation.is_scholarship = True
            allocation.scholarship_percentage = value
        else:
            allocation.discount_amount = value

        allocation.discount_reason = reason or allocation.discount_reason
        allocation.save()

        return allocation

    def _create_helpdesk_ticket(self, data, tenant):
        requester_type = data.get('requester_type')
        requester_id = data.get('requester_id')
        subject = data.get('subject')
        description = data.get('description')
        status_value = data.get('status') or 'OPEN'
        priority = data.get('priority') or 'MEDIUM'

        user = None
        if str(requester_type).upper() == 'STUDENT':
            student = Student.objects.filter(
                tenant=tenant,
                admission_number=requester_id,
                is_deleted=False
            ).first()
            user = student.user if student and hasattr(student, 'user') else None
        elif str(requester_type).upper() == 'STAFF':
            staff = Staff.objects.filter(
                tenant=tenant,
                employee_id=requester_id,
                is_deleted=False
            ).first()
            user = staff.user if staff and hasattr(staff, 'user') else None

        if not user:
            user = self._request_user

        ticket = HelpdeskTicket.objects.create(
            tenant=tenant,
            raised_by=user,
            category='OTHER',
            subject=subject,
            description=description,
            status=status_value,
            priority=priority
        )

        return ticket

    def _create_lms_course(self, data, tenant):
        course_name = data.get('course_name')
        course_code = data.get('course_code')
        class_name = data.get('class_name')
        subject_code = data.get('subject_code')
        teacher_employee_id = data.get('teacher_employee_id')
        start_time = self._parse_datetime(data.get('start_time'))
        duration_minutes = int(float(data.get('duration_minutes') or 45))
        meeting_link = data.get('meeting_link')

        grade_level = self._find_grade_level(tenant, class_name)
        if not grade_level:
            raise ValueError(f"Class '{class_name}' not found")

        subject = Subject.objects.filter(
            tenant=tenant,
            code__iexact=str(subject_code).strip()
        ).first()
        if not subject:
            raise ValueError(f"Subject '{subject_code}' not found")

        teacher = Staff.objects.filter(
            tenant=tenant,
            employee_id=teacher_employee_id,
            is_deleted=False
        ).first()
        if not teacher:
            raise ValueError(f"Teacher '{teacher_employee_id}' not found")

        if not start_time:
            start_date = DateParser.parse(data.get('start_date')) or date.today()
            start_time = datetime.combine(start_date, datetime.min.time())

        live_class = LiveClass.objects.create(
            tenant=tenant,
            title=course_name or course_code,
            description=data.get('description') or '',
            start_time=start_time,
            duration_minutes=duration_minutes,
            grade_level=grade_level,
            subject=subject,
            teacher=teacher,
            meeting_link=meeting_link or '',
        )

        return live_class

    def _create_lms_enrollment(self, data, tenant):
        raise ValueError("LMS enrollments are not supported (no enrollment model found)")

    def _create_idcard_record(self, data, tenant):
        card_number = data.get('card_number')
        admission_number = data.get('admission_number')
        employee_id = data.get('employee_id')
        template_code = data.get('template_code')
        issued_date = DateParser.parse(data.get('issued_date')) or date.today()
        expiry_date = DateParser.parse(data.get('expiry_date'))
        status_value = data.get('status') or 'active'
        file_url = data.get('file_url') or f"manual://{card_number}"

        entity_type = 'student' if admission_number else 'staff'
        entity = None
        if admission_number:
            entity = Student.objects.filter(
                tenant=tenant,
                admission_number=admission_number,
                is_deleted=False
            ).first()
        elif employee_id:
            entity = Staff.objects.filter(
                tenant=tenant,
                employee_id=employee_id,
                is_deleted=False
            ).first()

        if not entity:
            raise ValueError("Student or Staff not found for ID card")

        template = IDCardTemplate.objects.filter(
            tenant=tenant,
            name__iexact=str(template_code).strip(),
            is_active=True
        ).first()
        if not template:
            raise ValueError(f"ID Card Template '{template_code}' not found")

        valid_until = datetime.combine(expiry_date, datetime.min.time()) if expiry_date else datetime.now()

        record = IDCardRecord.objects.create(
            tenant=tenant,
            entity_type=entity_type,
            entity_id=entity.id,
            template=template,
            file_url=file_url,
            file_format='pdf',
            status=status_value,
            valid_until=valid_until
        )

        return record
    
    def _get_transformer(self, module):
        """Get data transformer for module"""
        return {
            'students': DataTransformer.transform_student,
            'staff': DataTransformer.transform_staff,
            'classes': DataTransformer.transform_class_section,
        }.get(module, lambda x: x)
    
    def _find_existing(self, module, field, value, tenant):
        """Find existing record by unique field"""
        models = {
            'students': Student,
            'staff': Staff,
            'subjects': Subject,
        }
        model = models.get(module)
        if model:
            try:
                return model.objects.filter(tenant=tenant, **{field: value}, is_deleted=False).first()
            except:
                pass
        return None
    
    def _create_record(self, module, data, tenant):
        """Create a new record based on module type"""
        creators = {
            'students': self._create_student,
            'staff': self._create_staff,
            'classes': self._create_class_section,
            'subjects': self._create_subject,
            'fee_structures': self._create_fee_structure,
            'fee_allocations': self._create_fee_allocation,
            'student_enrollments': self._create_enrollment,
            'parents': self._update_parent,
            'attendance': self._create_attendance,
            'transport': self._create_transport,
            'student_photos': self._create_student_photo,
            'exam_results': self._create_exam_result,
            'exam_schedule': self._create_exam_schedule,
            'timetable': self._create_timetable_slot,
            'library_books': self._create_library_book,
            'library_transactions': self._create_library_transaction,
            'payroll_payments': self._create_payroll_payment,
            'hostel_allocations': self._create_hostel_allocation,
            'inventory_items': self._create_inventory_item,
            'certificates_issued': self._create_certificate_issued,
            'finance_journal_entries': self._create_finance_journal_entry,
            'fee_discounts': self._apply_fee_discount,
            'helpdesk_tickets': self._create_helpdesk_ticket,
            'lms_courses': self._create_lms_course,
            'lms_enrollments': self._create_lms_enrollment,
            'idcards': self._create_idcard_record,
        }
        
        creator = creators.get(module)
        if creator:
            return creator(data, tenant)
        return None
    
    def _update_record(self, module, existing, data, tenant):
        """Update an existing record"""
        updaters = {
            'students': self._update_student,
            'staff': self._update_staff,
            'parents': self._update_parent,
        }
        
        updater = updaters.get(module)
        if updater:
            return updater(existing, data, tenant)
        return None
    
    # ========================================
    # RECORD CREATION METHODS
    # ========================================
    
    def _create_student(self, data, tenant):
        """Create a new student record"""
        student = Student.objects.create(
            tenant=tenant,
            admission_number=data.get('admission_number'),
            first_name=data.get('first_name', ''),
            middle_name=data.get('middle_name', ''),
            last_name=data.get('last_name', ''),
            date_of_birth=data.get('date_of_birth'),
            gender=data.get('gender', 'M'),
            blood_group=data.get('blood_group', ''),
            email=data.get('email', ''),
            phone=data.get('phone', ''),
            address=data.get('address', ''),
            father_name=data.get('father_name', ''),
            father_phone=data.get('father_phone', ''),
            father_email=data.get('father_email', ''),
            father_occupation=data.get('father_occupation', ''),
            mother_name=data.get('mother_name', ''),
            mother_phone=data.get('mother_phone', ''),
            mother_email=data.get('mother_email', ''),
            mother_occupation=data.get('mother_occupation', ''),
            guardian_name=data.get('guardian_name', ''),
            guardian_phone=data.get('guardian_phone', ''),
            guardian_relation=data.get('guardian_relation', ''),
            aadhar_number=data.get('aadhar_number', ''),
            pen_number=data.get('pen_number', ''),
            nationality=data.get('nationality', 'Indian'),
            religion=data.get('religion', ''),
            caste=data.get('caste', ''),
            admission_date=data.get('admission_date', date.today()),
            previous_school_name=data.get('previous_school_name', ''),
            transfer_certificate_number=data.get('transfer_certificate_number', ''),
            notes=data.get('notes', ''),
            is_active=True
        )
        
        # Create enrollment if class/section provided
        class_name = data.get('class_name')
        section_name = data.get('section_name')
        
        if class_name:
            self._create_student_enrollment(student, class_name, section_name, data, tenant)

        # Ensure parent portal accounts exist/linked for this student
        try:
            self._ensure_parent_accounts(student)
        except Exception as e:
            logger.warning(f"Failed to ensure parent accounts for student {student.id}: {e}")
        
        return student
    
    def _create_student_enrollment(self, student, class_name, section_name, data, tenant):
        """Create student enrollment in class/section"""
        try:
            grade_level = self._find_grade_level(tenant, class_name)
            
            if not grade_level:
                logger.warning(f"Grade level not found for value '{class_name}' (student {student.id})")
                return None
            
            section = None
            if section_name:
                section = self._find_section(tenant, grade_level, section_name)
                if not section:
                    logger.warning(
                        f"Section not found for value '{section_name}' in grade '{grade_level.name}' (student {student.id})"
                    )
            
            if not section:
                # Get first section of the grade if not specified
                section = Section.objects.filter(
                    tenant=tenant,
                    grade_level=grade_level,
                    is_deleted=False
                ).first()
            
            if section:
                academic_year = AcademicYear.objects.filter(
                    tenant=tenant,
                    is_active=True,
                    is_deleted=False
                ).first()

                if not academic_year:
                    academic_year = AcademicYear.objects.filter(
                        tenant=tenant,
                        is_deleted=False
                    ).order_by('-start_date').first()

                if not academic_year:
                    from datetime import timedelta
                    today = date.today()
                    try:
                        academic_year = AcademicYear.objects.create(
                            tenant=tenant,
                            name=f"{today.year}-{today.year + 1}",
                            start_date=today,
                            end_date=today + timedelta(days=365),
                            is_active=True,
                            is_enrollment_open=False,
                            is_locked=False,
                            description='Auto-created academic year during import'
                        )
                    except Exception as e:
                        logger.warning(f"Failed to create fallback academic year: {e}")
                        return None

                if academic_year:
                    enrollment, created = StudentEnrollment.objects.get_or_create(
                        tenant=tenant,
                        student=student,
                        academic_year=academic_year,
                        defaults={
                            'section': section,
                            'roll_number': data.get('roll_number', ''),
                            'status': 'ACTIVE',
                            'enrollment_date': data.get('admission_date', date.today())
                        }
                    )

                    # If enrollment already existed, update its section/roll/enrollment_date if provided
                    if not created:
                        updated = False
                        if section and enrollment.section != section:
                            enrollment.section = section
                            updated = True
                        roll = data.get('roll_number')
                        if roll and str(enrollment.roll_number) != str(roll):
                            enrollment.roll_number = roll
                            updated = True
                        enroll_date = data.get('admission_date')
                        if enroll_date and enrollment.enrollment_date != enroll_date:
                            enrollment.enrollment_date = enroll_date
                            updated = True
                        if updated:
                            enrollment.save()

                    return enrollment
        except Exception as e:
            logger.warning(f"Failed to create enrollment: {e}")
        
        return None

    def _ensure_parent_accounts(self, student):
        """Create or link parent portal accounts for a student based on parent fields."""
        from users.models import User
        from students.models import ParentUser

        # Ensure family_id exists for linking siblings
        if not student.family_id:
            student.family_id = f"FAM-{uuid.uuid4().hex[:8].upper()}"
            student.save(update_fields=['family_id'])

        # Link existing parent accounts by family_id if present
        existing_parents = ParentUser.objects.filter(
            tenant=student.tenant,
            students__family_id=student.family_id
        ).distinct()
        if existing_parents.exists():
            for parent in existing_parents:
                parent.students.add(student)
            return

        def _create_or_link(relation_type, name, phone, email):
            clean_phone = PhoneValidator.clean(phone) if phone else ''
            clean_email = EmailValidator.clean(email) if email else ''

            if not clean_phone and not clean_email:
                return

            user = None
            if clean_phone:
                user = User.objects.filter(tenant=student.tenant, phone_number=clean_phone).first()
            if not user and clean_email:
                user = User.objects.filter(email__iexact=clean_email).first()

            if user:
                parent_profile = ParentUser.objects.filter(user=user).first()
                if not parent_profile:
                    parent_profile = ParentUser.objects.create(
                        user=user,
                        tenant=student.tenant,
                        relation_type=relation_type,
                        portal_access_enabled=True
                    )
                parent_profile.students.add(student)
                return

            # Create new parent user
            import random
            import string
            temp_password = ''.join(random.choices(string.ascii_letters + string.digits, k=10))
            base_email = clean_email or (f"{clean_phone}@parent.local" if clean_phone else f"{student.id}@parent.local")

            # Ensure email uniqueness (email is globally unique)
            if User.objects.filter(email__iexact=base_email).exists():
                base_email = f"{base_email.split('@')[0]}.{student.tenant_id}@parent.local"

            # Name split
            first_name = 'Parent'
            last_name = ''
            if name:
                parts = str(name).strip().split()
                if parts:
                    first_name = parts[0]
                    last_name = ' '.join(parts[1:]) if len(parts) > 1 else ''

            user = User.objects.create_user(
                tenant=student.tenant,
                email=base_email,
                phone_number=clean_phone,
                first_name=first_name,
                last_name=last_name,
                password=temp_password
            )

            parent_profile = ParentUser.objects.create(
                user=user,
                tenant=student.tenant,
                relation_type=relation_type,
                portal_access_enabled=True
            )
            parent_profile.students.add(student)

        _create_or_link('FATHER', student.father_name, student.father_phone, student.father_email)
        _create_or_link('MOTHER', student.mother_name, student.mother_phone, student.mother_email)
        _create_or_link('GUARDIAN', student.guardian_name, student.guardian_phone, None)
    
    def _update_student(self, student, data, tenant):
        """Update existing student record"""
        update_fields = [
            'first_name', 'middle_name', 'last_name', 'date_of_birth', 'gender',
            'blood_group', 'email', 'phone', 'address', 'father_name', 'father_phone',
            'father_email', 'father_occupation', 'mother_name', 'mother_phone',
            'mother_email', 'mother_occupation', 'guardian_name', 'guardian_phone',
            'guardian_relation', 'aadhar_number', 'pen_number', 'nationality',
            'religion', 'caste', 'notes'
        ]
        
        for field in update_fields:
            if field in data and data[field]:
                setattr(student, field, data[field])
        
        student.save()

        # Update enrollment/class-section if provided in import data
        class_name = data.get('class_name') or data.get('grade') or data.get('grade_level')
        section_name = data.get('section_name') or data.get('section')
        if class_name:
            try:
                # This will create or update the student's enrollment for the active academic year
                self._create_student_enrollment(student, class_name, section_name, data, tenant)
            except Exception as e:
                logger.warning(f"Failed to update enrollment for student {student.id}: {e}")

        # Ensure parent portal accounts exist/linked for this student
        try:
            self._ensure_parent_accounts(student)
        except Exception as e:
            logger.warning(f"Failed to ensure parent accounts for student {student.id}: {e}")

        return student
    
    def _create_staff(self, data, tenant):
        """Create a new staff record"""
        # Map gender from import format to model format
        gender = data.get('gender', '')
        if gender == 'M':
            gender = 'MALE'
        elif gender == 'F':
            gender = 'FEMALE'
        elif gender == 'O':
            gender = 'OTHER'
        
        staff = Staff.objects.create(
            tenant=tenant,
            employee_id=data.get('employee_id'),
            first_name=data.get('first_name', ''),
            middle_name=data.get('middle_name', ''),
            last_name=data.get('last_name', ''),
            date_of_birth=data.get('date_of_birth'),
            gender=gender,
            blood_group=data.get('blood_group', ''),
            email=data.get('email', ''),
            phone=data.get('phone', ''),
            alternate_phone=data.get('alternate_phone', ''),
            address=data.get('address', ''),
            city=data.get('city', ''),
            state=data.get('state', ''),
            postal_code=data.get('postal_code', ''),
            joining_date=data.get('joining_date') or date.today(),
            designation=data.get('designation', 'TEACHER'),
            employment_type=data.get('employment_type', 'PERMANENT'),
            aadhar_number=data.get('aadhar_number', ''),
            pan_number=data.get('pan_number', ''),
            salary=data.get('salary'),
            bank_account_number=data.get('bank_account_number', ''),
            bank_name=data.get('bank_name', ''),
            bank_ifsc=data.get('bank_ifsc', ''),
            emergency_contact_name=data.get('emergency_contact_name', ''),
            emergency_contact_phone=data.get('emergency_contact_phone', ''),
            experience_years=data.get('experience_years', 0),
            status='ACTIVE'
        )
        
        # Link to department if provided
        if data.get('department'):
            dept_name = (data.get('department') or '').strip()
            if dept_name:
                # Normalize and get or create department to avoid duplicate errors
                dept, _ = Department.objects.get_or_create(
                    tenant=tenant,
                    name__iexact=dept_name,
                    defaults={
                        'name': dept_name,
                        'is_active': True
                    }
                )
                staff.department = dept
                staff.save()
        
        return staff
    
    def _update_staff(self, staff, data, tenant):
        """Update existing staff record"""
        update_fields = [
            'first_name', 'middle_name', 'last_name', 'date_of_birth', 'email',
            'phone', 'alternate_phone', 'address', 'city', 'state', 'postal_code',
            'designation', 'employment_type', 'aadhar_number', 'pan_number',
            'bank_account_number', 'bank_name', 'bank_ifsc',
            'emergency_contact_name', 'emergency_contact_phone'
        ]
        
        for field in update_fields:
            if field in data and data[field]:
                setattr(staff, field, data[field])
        
        staff.save()
        return staff
    
    def _create_class_section(self, data, tenant):
        """Create class and section"""
        raw_class = data.get('class_name', '')
        raw_section = data.get('section_name', '')
        class_name, short_name = self._canonical_grade_fields(raw_class)
        section_name = self._normalize_section_name(raw_section) or str(raw_section).strip()
        
        if not class_name or not section_name:
            return None
        
        # Get or create grade level. Compute a safe display_order to avoid unique constraint collisions
        try:
            max_order = GradeLevel.objects.filter(tenant=tenant, is_deleted=False).aggregate(max_order=Max('display_order'))['max_order']
        except Exception:
            max_order = None

        next_order = (max_order or 0) + 1

        grade_level, created = GradeLevel.objects.get_or_create(
            tenant=tenant,
            name__iexact=class_name,
            defaults={
                'name': class_name,
                'short_name': short_name,
                'display_order': next_order,
                'is_active': True
            }
        )
        
        # Create section
        section, created = Section.objects.get_or_create(
            tenant=tenant,
            grade_level=grade_level,
            name__iexact=section_name,
            defaults={
                'name': section_name,
                'room_number': data.get('room_number', ''),
                'capacity': data.get('capacity', 40),
                'is_active': True
            }
        )
        
        return section
    
    def _create_subject(self, data, tenant):
        """Create subject"""
        raw_name = (data.get('subject_name') or '').strip()
        if not raw_name:
            return None

        # Normalize name (collapse whitespace)
        norm_name = ' '.join(raw_name.split())

        # Use case-insensitive lookup to avoid duplicates, create if missing
        subject = Subject.objects.filter(
            tenant=tenant,
            name__iexact=norm_name,
            is_deleted=False
        ).first()

        if subject:
            return subject

        # Create new subject with provided metadata
        subject = Subject.objects.create(
            tenant=tenant,
            name=norm_name,
            code=(data.get('subject_code') or '').strip(),
            description=(data.get('description') or '').strip(),
            is_active=True
        )
        return subject
    
    def _create_fee_structure(self, data, tenant):
        """Create fee structure"""
        fee_type_name = data.get('fee_type', '').strip()
        class_name = data.get('class_name', '').strip()
        
        # Get or create fee category
        category, _ = FeeCategory.objects.get_or_create(
            tenant=tenant,
            name__iexact=fee_type_name,
            defaults={
                'name': fee_type_name,
                'code': fee_type_name[:10].upper().replace(' ', '_'),
                'is_active': True
            }
        )
        
        # Get academic year: prefer active, else latest; if none, create a default one
        academic_year = AcademicYear.objects.filter(
            tenant=tenant,
            is_active=True,
            is_deleted=False
        ).first()

        if not academic_year:
            academic_year = AcademicYear.objects.filter(
                tenant=tenant,
                is_deleted=False
            ).order_by('-start_date').first()

        if not academic_year:
            # Create a fallback academic year spanning today -> today+1year
            from datetime import timedelta
            today = date.today()
            try:
                new_name = f"{today.year}-{today.year + 1}"
                academic_year = AcademicYear.objects.create(
                    tenant=tenant,
                    name=new_name,
                    start_date=today,
                    end_date=today + timedelta(days=365),
                    is_active=True,
                    is_enrollment_open=False,
                    is_locked=False,
                    description='Auto-created academic year during import'
                )
            except Exception as e:
                raise ValueError(f"No active academic year found and failed to create fallback: {e}")
        
        # Create or get fee structure to avoid duplicates
        fs_defaults = {
            'amount': data.get('amount', 0),
            'frequency': data.get('frequency', 'MONTHLY'),
            'due_day': data.get('due_day', 10),
            'is_mandatory': data.get('is_mandatory', 'true').lower() == 'true' if isinstance(data.get('is_mandatory'), str) else True,
            'is_active': True
        }

        fee_structure, created = FeeStructure.objects.get_or_create(
            tenant=tenant,
            academic_year=academic_year,
            category=category,
            class_level=class_name,
            defaults=fs_defaults
        )

        # If existing but different values, optionally update - keep existing values for now
        return fee_structure
    
    def _create_enrollment(self, data, tenant):
        """Create student enrollment"""
        admission_number = data.get('admission_number', '').strip()
        
        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        
        if not student:
            raise ValueError(f"Student not found: {admission_number}")
        
        return self._create_student_enrollment(
            student, 
            data.get('class_name'), 
            data.get('section_name'),
            data,
            tenant
        )
    
    def _update_parent(self, data_or_student, data_or_tenant, tenant=None):
        """Update parent information for existing student"""
        # Handle both update scenarios
        if isinstance(data_or_student, Student):
            student = data_or_student
            data = data_or_tenant
        else:
            data = data_or_student
            tenant = data_or_tenant
            admission_number = data.get('admission_number', '').strip()
            
            student = Student.objects.filter(
                tenant=tenant,
                admission_number=admission_number,
                is_deleted=False
            ).first()
            
            if not student:
                raise ValueError(f"Student not found: {admission_number}")
        
        # Update parent fields
        parent_fields = [
            'father_name', 'father_phone', 'father_email', 'father_occupation',
            'mother_name', 'mother_phone', 'mother_email', 'mother_occupation',
            'guardian_name', 'guardian_phone', 'guardian_relation'
        ]
        
        for field in parent_fields:
            if field in data and data[field]:
                setattr(student, field, data[field])
        
        student.save()
        return student
    
    def _create_attendance(self, data, tenant):
        """Create attendance record"""
        admission_number = data.get('admission_number', '').strip()
        
        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        
        if not student:
            raise ValueError(f"Student not found: {admission_number}")
        
        academic_year = AcademicYear.objects.filter(
            tenant=tenant,
            is_active=True,
            is_deleted=False
        ).first()
        
        attendance_date = DateParser.parse(data.get('date'))
        if not attendance_date:
            raise ValueError("Invalid date")
        
        record, created = AttendanceRecord.objects.get_or_create(
            tenant=tenant,
            student=student,
            date=attendance_date,
            defaults={
                'academic_year': academic_year,
                'record_type': 'STUDENT',
                'status': data.get('status', 'PRESENT'),
                'remarks': data.get('remarks', ''),
                'marking_method': 'MANUAL',
            }
        )
        
        return record
    
    def _create_transport(self, data, tenant):
        """Create transport allocation"""
        admission_number = data.get('admission_number', '').strip()
        route_name = data.get('route_name', '').strip()
        stop_name = data.get('stop_name', '').strip()
        
        student = Student.objects.filter(
            tenant=tenant,
            admission_number=admission_number,
            is_deleted=False
        ).first()
        
        if not student:
            raise ValueError(f"Student not found: {admission_number}")
        
        route = Route.objects.filter(
            tenant=tenant,
            name__iexact=route_name,
            is_deleted=False
        ).first()
        
        if not route:
            raise ValueError(f"Route not found: {route_name}")
        
        stop = Stop.objects.filter(
            tenant=tenant,
            route=route,
            name__iexact=stop_name,
            is_deleted=False
        ).first()
        
        if not stop:
            raise ValueError(f"Stop not found: {stop_name}")
        
        transport, created = StudentTransport.objects.update_or_create(
            tenant=tenant,
            student=student,
            defaults={
                'stop': stop,
                'is_active': True,
                'start_date': DateParser.parse(data.get('start_date'), date.today())
            }
        )
        
        return transport


class ImportJobStatusView(APIView):
    """Check status of an import job"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request, job_id):
        """Get import job status"""
        try:
            job = ImportJob.objects.get(
                id=job_id,
                tenant=request.user.tenant
            )
            
            return Response({
                'id': str(job.id),
                'module': job.module,
                'status': job.status,
                'total_rows': job.total_rows,
                'processed_rows': job.processed_rows,
                'successful_rows': job.successful_rows,
                'failed_rows': job.failed_rows,
                'duplicate_rows': job.duplicate_rows,
                'errors': job.validation_errors[:20],
                'created_at': job.created_at.isoformat(),
                'completed_at': job.completed_at.isoformat() if job.completed_at else None,
            })
        except ImportJob.DoesNotExist:
            return Response(
                {'error': 'Import job not found'},
                status=status.HTTP_404_NOT_FOUND
            )


class ImportJobHistoryView(APIView):
    """List recent import jobs"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request):
        """Get import job history"""
        jobs = ImportJob.objects.filter(
            tenant=request.user.tenant
        ).order_by('-created_at')[:50]
        
        return Response({
            'jobs': [
                {
                    'id': str(job.id),
                    'module': job.module,
                    'status': job.status,
                    'filename': job.original_filename,
                    'total_rows': job.total_rows,
                    'successful_rows': job.successful_rows,
                    'failed_rows': job.failed_rows,
                    'created_at': job.created_at.isoformat(),
                    'created_by': job.created_by.email if job.created_by else None,
                }
                for job in jobs
            ]
        })


class RollbackImportView(APIView):
    """Rollback a completed import"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def post(self, request, job_id):
        """Rollback an import job"""
        try:
            job = ImportJob.objects.get(
                id=job_id,
                tenant=request.user.tenant,
                status='COMPLETED'
            )
            
            if not job.created_record_ids:
                return Response({
                    'success': False,
                    'error': 'No records to rollback'
                })
            
            # Get the model for this module
            models_map = {
                'students': Student,
                'staff': Staff,
                'classes': Section,
                'subjects': Subject,
            }
            
            model = models_map.get(job.module)
            if not model:
                return Response({
                    'success': False,
                    'error': 'Rollback not supported for this module'
                })
            
            # Delete created records
            deleted_count = model.objects.filter(
                id__in=job.created_record_ids,
                tenant=request.user.tenant
            ).delete()[0]
            
            job.status = 'ROLLED_BACK'
            job.save()
            
            return Response({
                'success': True,
                'deleted_count': deleted_count
            })
            
        except ImportJob.DoesNotExist:
            return Response(
                {'error': 'Import job not found or cannot be rolled back'},
                status=status.HTTP_404_NOT_FOUND
            )


class ExportDataView(APIView):
    """Export data to Excel/CSV"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request, module):
        """Export module data"""
        format_type = request.query_params.get('format', 'xlsx')
        
        template = get_template(module)
        if not template:
            return Response(
                {'error': f'Invalid module: {module}'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        tenant = request.user.tenant
        
        # Get data based on module
        data = self._get_export_data(module, tenant)
        
        if format_type == 'xlsx' and OPENPYXL_AVAILABLE:
            return self._export_xlsx(module, template, data)
        else:
            return self._export_csv(module, template, data)
    
    def _get_export_data(self, module, tenant):
        """Get data for export"""
        if module == 'students':
            students = Student.objects.filter(tenant=tenant, is_deleted=False)
            return [self._student_to_dict(s) for s in students]
        elif module == 'staff':
            staff = Staff.objects.filter(tenant=tenant, is_deleted=False)
            return [self._staff_to_dict(s) for s in staff]
        elif module == 'classes':
            sections = Section.objects.filter(tenant=tenant, is_deleted=False)
            return [self._section_to_dict(s) for s in sections]
        elif module == 'fee_allocations':
            allocations = FeeAllocation.objects.filter(tenant=tenant, is_deleted=False)
            return [self._fee_allocation_to_dict(a) for a in allocations]
        elif module == 'student_photos':
            students = Student.objects.filter(tenant=tenant, is_deleted=False).exclude(photo='')
            return [self._student_photo_to_dict(s) for s in students]
        elif module == 'exam_results':
            results = ExamResult.objects.filter(tenant=tenant, is_deleted=False)
            return [self._exam_result_to_dict(r) for r in results]
        elif module == 'exam_schedule':
            schedules = ExamSchedule.objects.filter(tenant=tenant, is_deleted=False)
            return [self._exam_schedule_to_dict(s) for s in schedules]
        elif module == 'timetable':
            slots = TimetableSlot.objects.filter(tenant=tenant, is_deleted=False)
            return [self._timetable_to_dict(s) for s in slots]
        elif module == 'library_books':
            books = Book.objects.filter(tenant=tenant, is_deleted=False)
            return [self._library_book_to_dict(b) for b in books]
        elif module == 'library_transactions':
            issues = BookIssue.objects.filter(tenant=tenant, is_deleted=False)
            return [self._library_issue_to_dict(i) for i in issues]
        elif module == 'payroll_payments':
            payments = SalaryPayment.objects.filter(tenant=tenant, is_deleted=False)
            return [self._salary_payment_to_dict(p) for p in payments]
        elif module == 'hostel_allocations':
            allocations = HostelAllocation.objects.filter(tenant=tenant, is_deleted=False)
            return [self._hostel_allocation_to_dict(a) for a in allocations]
        elif module == 'inventory_items':
            items = Item.objects.filter(tenant=tenant, is_deleted=False)
            return [self._inventory_item_to_dict(i) for i in items]
        elif module == 'certificates_issued':
            certs = GeneratedCertificate.objects.filter(tenant=tenant, is_deleted=False)
            return [self._certificate_to_dict(c) for c in certs]
        elif module == 'finance_journal_entries':
            entries = JournalEntry.objects.filter(tenant=tenant)
            return self._journal_entries_to_dict(entries)
        elif module == 'fee_discounts':
            allocations = FeeAllocation.objects.filter(tenant=tenant, is_deleted=False)
            return [self._fee_discount_to_dict(a) for a in allocations]
        elif module == 'helpdesk_tickets':
            tickets = HelpdeskTicket.objects.filter(tenant=tenant, is_deleted=False)
            return [self._helpdesk_ticket_to_dict(t) for t in tickets]
        elif module == 'lms_courses':
            live_classes = LiveClass.objects.filter(tenant=tenant, is_deleted=False)
            return [self._live_class_to_dict(c) for c in live_classes]
        elif module == 'lms_enrollments':
            return []
        elif module == 'idcards':
            records = IDCardRecord.objects.filter(tenant=tenant, is_deleted=False)
            return [self._idcard_to_dict(r) for r in records]
        return []
    
    def _student_to_dict(self, student):
        """Convert student to dictionary for export"""
        enrollment = student.enrollments.filter(status='ACTIVE').first()
        
        return {
            'admission_number': student.admission_number,
            'first_name': student.first_name,
            'middle_name': student.middle_name,
            'last_name': student.last_name,
            'date_of_birth': DateParser.format_for_display(student.date_of_birth) if student.date_of_birth else '',
            'gender': student.gender,
            'blood_group': student.blood_group,
            'email': student.email,
            'phone': student.phone,
            'address': student.address,
            'father_name': student.father_name,
            'father_phone': student.father_phone,
            'father_email': student.father_email,
            'mother_name': student.mother_name,
            'mother_phone': student.mother_phone,
            'mother_email': student.mother_email,
            'aadhar_number': student.aadhar_number,
            'class_name': enrollment.section.grade_level.name if enrollment and enrollment.section else '',
            'section_name': enrollment.section.name if enrollment and enrollment.section else '',
            'roll_number': enrollment.roll_number if enrollment else '',
        }
    
    def _staff_to_dict(self, staff):
        """Convert staff to dictionary for export"""
        return {
            'employee_id': staff.employee_id,
            'first_name': staff.first_name,
            'middle_name': staff.middle_name,
            'last_name': staff.last_name,
            'date_of_birth': DateParser.format_for_display(staff.date_of_birth) if staff.date_of_birth else '',
            'gender': staff.gender,
            'email': staff.email,
            'phone': staff.phone,
            'designation': staff.designation,
            'department': staff.department.name if staff.department else '',
            'joining_date': DateParser.format_for_display(staff.joining_date) if staff.joining_date else '',
        }
    
    def _section_to_dict(self, section):
        """Convert section to dictionary for export"""
        return {
            'class_name': section.grade_level.name,
            'section_name': section.name,
            'room_number': section.room_number,
            'capacity': section.capacity,
        }

    def _fee_allocation_to_dict(self, allocation):
        return {
            'admission_number': allocation.student.admission_number if allocation.student else '',
            'fee_type': allocation.fee_structure.category.name if allocation.fee_structure else '',
            'amount': str(allocation.custom_amount or allocation.fee_structure.amount) if allocation.fee_structure else '',
            'frequency': allocation.fee_structure.frequency if allocation.fee_structure else '',
            'academic_year': allocation.fee_structure.academic_year.name if allocation.fee_structure else '',
        }

    def _student_photo_to_dict(self, student):
        return {
            'admission_number': student.admission_number,
            'photo_filename': getattr(student.photo, 'name', '') if student.photo else '',
            'photo_url': getattr(student.photo, 'url', '') if student.photo else '',
        }

    def _exam_result_to_dict(self, result):
        return {
            'admission_number': result.student.admission_number if result.student else '',
            'exam_name': result.exam.name if result.exam else '',
            'subject_code': result.exam.subject.code if result.exam and result.exam.subject else '',
            'max_marks': str(result.exam.total_marks) if result.exam else '',
            'marks_obtained': str(result.marks_obtained),
            'grade': result.grade,
            'exam_date': DateParser.format_for_display(result.exam.exam_term.start_date) if result.exam and result.exam.exam_term else '',
            'academic_year': result.exam.exam_term.academic_year.name if result.exam and result.exam.exam_term else '',
        }

    def _exam_schedule_to_dict(self, schedule):
        return {
            'exam_name': schedule.exam.name if schedule.exam else '',
            'class_name': schedule.exam.grade_level.name if schedule.exam and schedule.exam.grade_level else '',
            'section_name': schedule.section.name if schedule.section else '',
            'subject_code': schedule.exam.subject.code if schedule.exam and schedule.exam.subject else '',
            'exam_date': DateParser.format_for_display(schedule.exam_date),
            'start_time': schedule.start_time.strftime('%H:%M') if schedule.start_time else '',
            'end_time': schedule.end_time.strftime('%H:%M') if schedule.end_time else '',
            'max_marks': str(schedule.exam.total_marks) if schedule.exam else '',
            'academic_year': schedule.exam.exam_term.academic_year.name if schedule.exam and schedule.exam.exam_term else '',
        }

    def _timetable_to_dict(self, slot):
        return {
            'class_name': slot.section.grade_level.name if slot.section else '',
            'section_name': slot.section.name if slot.section else '',
            'academic_year': slot.academic_year.name if slot.academic_year else '',
            'day_of_week': slot.day_of_week,
            'start_time': slot.start_time.strftime('%H:%M') if slot.start_time else '',
            'end_time': slot.end_time.strftime('%H:%M') if slot.end_time else '',
            'period_number': slot.period_number or '',
            'subject_code': slot.subject.code if slot.subject else '',
            'teacher_employee_id': slot.teacher.employee_id if slot.teacher else '',
            'room_number': slot.room,
        }

    def _library_book_to_dict(self, book):
        return {
            'isbn': book.isbn,
            'title': book.title,
            'author': book.author,
            'publisher': book.publisher,
            'category': book.category,
            'copies_total': str(book.total_copies),
            'location': book.shelf_location,
            'published_year': book.year_published or '',
        }

    def _library_issue_to_dict(self, issue):
        return {
            'transaction_id': f"LIB{issue.id}",
            'admission_number': issue.member.student.admission_number if issue.member and issue.member.student else '',
            'isbn': issue.copy.book.isbn if issue.copy and issue.copy.book else '',
            'issue_date': issue.issued_date.date().isoformat() if issue.issued_date else '',
            'due_date': issue.due_date.date().isoformat() if issue.due_date else '',
            'return_date': issue.returned_date.date().isoformat() if issue.returned_date else '',
            'fine_amount': str(issue.fine_amount),
            'status': issue.status,
        }

    def _salary_payment_to_dict(self, payment):
        return {
            'payment_reference': payment.payment_number,
            'employee_id': payment.staff.employee_id if payment.staff else '',
            'period_start': payment.month.isoformat() if payment.month else '',
            'period_end': payment.month.isoformat() if payment.month else '',
            'gross_amount': str(payment.basic_salary + payment.allowances),
            'deductions': str(payment.deductions),
            'net_amount': str(payment.net_salary),
            'payment_date': payment.payment_date.isoformat() if payment.payment_date else '',
            'payment_mode': '',
            'remarks': payment.remarks,
        }

    def _hostel_allocation_to_dict(self, allocation):
        return {
            'admission_number': allocation.student.admission_number if allocation.student else '',
            'hostel_name': allocation.bed.room.building.name if allocation.bed else '',
            'room_number': allocation.bed.room.room_number if allocation.bed else '',
            'bed_number': allocation.bed.bed_number if allocation.bed else '',
            'start_date': DateParser.format_for_display(allocation.start_date) if allocation.start_date else '',
            'end_date': DateParser.format_for_display(allocation.end_date) if allocation.end_date else '',
            'status': 'ACTIVE' if allocation.is_active else 'INACTIVE',
        }

    def _inventory_item_to_dict(self, item):
        return {
            'item_code': item.sku,
            'item_name': item.name,
            'category': item.category.name if item.category else '',
            'quantity': str(item.current_stock),
            'unit_cost': str(item.cost_price),
            'purchase_date': '',
            'vendor': '',
            'location': '',
            'condition': '',
        }

    def _certificate_to_dict(self, cert):
        return {
            'certificate_number': cert.certificate_number,
            'admission_number': cert.request.student.admission_number if cert.request and cert.request.student else '',
            'certificate_type': cert.request.template.name if cert.request and cert.request.template else '',
            'issue_date': DateParser.format_for_display(cert.issued_date) if cert.issued_date else '',
            'expiry_date': '',
            'remarks': cert.request.reason if cert.request else '',
        }

    def _journal_entries_to_dict(self, entries):
        rows = []
        for entry in entries:
            for line in entry.lines.all():
                rows.append({
                    'entry_number': entry.entry_number,
                    'entry_date': entry.entry_date.isoformat() if entry.entry_date else '',
                    'account_code': line.account.code if line.account else '',
                    'account_name': line.account.name if line.account else '',
                    'debit': str(line.debit_amount),
                    'credit': str(line.credit_amount),
                    'narration': line.description or entry.description,
                    'reference': entry.reference_type or '',
                })
        return rows

    def _fee_discount_to_dict(self, allocation):
        discount_type = 'PERCENT' if allocation.scholarship_percentage > 0 else 'AMOUNT'
        value = allocation.scholarship_percentage if allocation.scholarship_percentage > 0 else allocation.discount_amount
        return {
            'admission_number': allocation.student.admission_number if allocation.student else '',
            'discount_code': allocation.discount_reason or 'DISCOUNT',
            'discount_type': discount_type,
            'value': str(value),
            'start_date': '',
            'end_date': '',
            'reason': allocation.discount_reason,
        }

    def _helpdesk_ticket_to_dict(self, ticket):
        return {
            'ticket_number': f"TICKET-{ticket.id}",
            'requester_type': 'STAFF' if ticket.raised_by and getattr(ticket.raised_by, 'is_staff', False) else 'STUDENT',
            'requester_id': ticket.raised_by.email if ticket.raised_by else '',
            'subject': ticket.subject,
            'description': ticket.description,
            'status': ticket.status,
            'priority': ticket.priority,
            'created_at': ticket.created_at.isoformat() if ticket.created_at else '',
            'resolved_at': ticket.resolved_at.isoformat() if ticket.resolved_at else '',
        }

    def _live_class_to_dict(self, live_class):
        return {
            'course_code': live_class.id,
            'class_name': live_class.grade_level.name if live_class.grade_level else '',
            'subject_code': live_class.subject.code if live_class.subject else '',
            'course_name': live_class.title,
            'description': live_class.description,
            'teacher_employee_id': live_class.teacher.employee_id if live_class.teacher else '',
            'start_time': live_class.start_time.isoformat() if live_class.start_time else '',
            'duration_minutes': live_class.duration_minutes,
            'meeting_link': live_class.meeting_link,
            'start_date': live_class.start_time.date().isoformat() if live_class.start_time else '',
            'end_date': '',
            'status': 'ACTIVE' if not live_class.is_completed else 'INACTIVE',
        }

    def _idcard_to_dict(self, record):
        admission_number = ''
        employee_id = ''
        if record.entity_type == 'student':
            student = Student.objects.filter(id=record.entity_id, is_deleted=False).first()
            admission_number = student.admission_number if student else ''
        if record.entity_type == 'staff':
            staff = Staff.objects.filter(id=record.entity_id, is_deleted=False).first()
            employee_id = staff.employee_id if staff else ''

        return {
            'card_number': record.id,
            'admission_number': admission_number,
            'employee_id': employee_id,
            'template_code': record.template.name if record.template else '',
            'file_url': record.file_url,
            'issued_date': record.issued_date.date().isoformat() if record.issued_date else '',
            'expiry_date': record.valid_until.date().isoformat() if record.valid_until else '',
            'status': record.status,
            'qr_code': '',
        }
    
    def _export_xlsx(self, module, template, data):
        """Export to Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = template.display_name
        
        # Headers
        field_names = [f.name for f in template.fields]
        headers = [f.display_name for f in template.fields]
        
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='E5E7EB', fill_type='solid')
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        # Data rows
        for row_num, row_data in enumerate(data, 2):
            for col, field_name in enumerate(field_names, 1):
                value = row_data.get(field_name, '')
                ws.cell(row=row_num, column=col, value=value)
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{module}_export_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response
    
    def _export_csv(self, module, template, data):
        """Export to CSV"""
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{module}_export_{datetime.now().strftime("%Y%m%d")}.csv"'
        
        response.write('\ufeff')  # BOM for Excel
        
        writer = csv.writer(response)
        
        field_names = [f.name for f in template.fields]
        headers = [f.display_name for f in template.fields]
        
        writer.writerow(headers)
        
        for row_data in data:
            row = [row_data.get(f, '') for f in field_names]
            writer.writerow(row)
        
        return response


class FullBackupView(APIView):
    """Create full backup of all tenant data"""
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    
    def get(self, request):
        """Create and download full backup"""
        tenant = request.user.tenant
        
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for module_name, template in TEMPLATE_REGISTRY.items():
                try:
                    # Get export view's data method
                    export_view = ExportDataView()
                    data = export_view._get_export_data(module_name, tenant)
                    
                    if not data:
                        continue
                    
                    # Create CSV
                    csv_buffer = io.StringIO()
                    writer = csv.DictWriter(csv_buffer, fieldnames=[f.name for f in template.fields])
                    writer.writeheader()
                    
                    for row in data:
                        writer.writerow({k: v for k, v in row.items() if k in [f.name for f in template.fields]})
                    
                    zip_file.writestr(f'{module_name}.csv', csv_buffer.getvalue())
                    
                except Exception as e:
                    logger.warning(f"Error exporting {module_name}: {e}")
            
            # Add metadata
            metadata = {
                'tenant_id': str(tenant.id),
                'tenant_name': tenant.name,
                'backup_date': datetime.now().isoformat(),
                'modules': list(TEMPLATE_REGISTRY.keys()),
            }
            zip_file.writestr('backup_metadata.json', json.dumps(metadata, indent=2))
        
        zip_buffer.seek(0)
        
        response = HttpResponse(zip_buffer.read(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="school_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip"'
        return response
