"""
Sample Data Generator for Testing Data Migration
Creates sample CSV/Excel files for testing import functionality
"""

import csv
import io
from datetime import date, timedelta
import random

# Try to import openpyxl
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class SampleDataGenerator:
    """Generate sample data for testing imports"""
    
    # Sample data pools
    FIRST_NAMES_MALE = ['Rahul', 'Amit', 'Pradeep', 'Vikram', 'Suresh', 'Rajesh', 'Arun', 'Vijay', 'Sanjay', 'Deepak']
    FIRST_NAMES_FEMALE = ['Priya', 'Anjali', 'Sneha', 'Pooja', 'Divya', 'Neha', 'Swati', 'Kavita', 'Meera', 'Asha']
    LAST_NAMES = ['Sharma', 'Kumar', 'Singh', 'Verma', 'Gupta', 'Patel', 'Jain', 'Mehta', 'Shah', 'Rao']
    
    CITIES = ['Delhi', 'Mumbai', 'Bangalore', 'Hyderabad', 'Chennai', 'Kolkata', 'Pune', 'Jaipur']
    RELIGIONS = ['Hindu', 'Muslim', 'Christian', 'Sikh', 'Buddhist', 'Jain']
    BLOOD_GROUPS = ['A+', 'A-', 'B+', 'B-', 'AB+', 'AB-', 'O+', 'O-']
    
    DESIGNATIONS = ['TEACHER', 'TEACHER', 'TEACHER', 'ASSISTANT_TEACHER', 'LIBRARIAN', 'ACCOUNTANT', 'CLERK']
    DEPARTMENTS = ['Mathematics', 'Science', 'English', 'Hindi', 'Social Studies', 'Computer Science', 'Physical Education']
    
    @classmethod
    def generate_students(cls, count=50, start_admission=1):
        """Generate sample student data"""
        students = []
        
        for i in range(count):
            gender = random.choice(['M', 'F'])
            first_name = random.choice(cls.FIRST_NAMES_MALE if gender == 'M' else cls.FIRST_NAMES_FEMALE)
            last_name = random.choice(cls.LAST_NAMES)
            
            # Random date of birth (between 6 and 18 years ago)
            age_days = random.randint(6*365, 18*365)
            dob = date.today() - timedelta(days=age_days)
            
            # Father and mother details
            father_first = random.choice(cls.FIRST_NAMES_MALE)
            mother_first = random.choice(cls.FIRST_NAMES_FEMALE)
            
            student = {
                'first_name': first_name,
                'last_name': last_name,
                'admission_number': f'STU{date.today().year}{start_admission + i:04d}',
                'date_of_birth': dob.strftime('%d-%m-%Y'),
                'gender': gender,
                'blood_group': random.choice(cls.BLOOD_GROUPS),
                'email': f'{first_name.lower()}.{last_name.lower()}{i}@example.com',
                'phone': f'9{random.randint(100000000, 999999999)}',
                'address': f'{random.randint(1, 500)}, {random.choice(["Main Street", "Park Road", "Gandhi Nagar", "Civil Lines"])}, {random.choice(cls.CITIES)}',
                'father_name': f'{father_first} {last_name}',
                'father_phone': f'9{random.randint(100000000, 999999999)}',
                'father_email': f'{father_first.lower()}.{last_name.lower()}@example.com',
                'father_occupation': random.choice(['Business', 'Service', 'Doctor', 'Engineer', 'Teacher', 'Farmer']),
                'mother_name': f'{mother_first} {last_name}',
                'mother_phone': f'9{random.randint(100000000, 999999999)}',
                'mother_occupation': random.choice(['Homemaker', 'Teacher', 'Doctor', 'Business', 'Service']),
                'nationality': 'Indian',
                'religion': random.choice(cls.RELIGIONS),
                'class_name': f'Class {random.randint(1, 12)}',
                'section_name': random.choice(['A', 'B', 'C']),
                'roll_number': str(random.randint(1, 60)),
            }
            students.append(student)
        
        return students
    
    @classmethod
    def generate_staff(cls, count=20, start_employee=1):
        """Generate sample staff data"""
        staff = []
        
        for i in range(count):
            gender = random.choice(['M', 'F'])
            first_name = random.choice(cls.FIRST_NAMES_MALE if gender == 'M' else cls.FIRST_NAMES_FEMALE)
            last_name = random.choice(cls.LAST_NAMES)
            
            # Random date of birth (between 25 and 60 years ago)
            age_days = random.randint(25*365, 60*365)
            dob = date.today() - timedelta(days=age_days)
            
            # Joining date (1-15 years ago)
            join_days = random.randint(365, 15*365)
            join_date = date.today() - timedelta(days=join_days)
            
            staff_member = {
                'first_name': first_name,
                'last_name': last_name,
                'employee_id': f'EMP{start_employee + i:04d}',
                'email': f'{first_name.lower()}.{last_name.lower()}@school.edu',
                'phone': f'9{random.randint(100000000, 999999999)}',
                'date_of_birth': dob.strftime('%d-%m-%Y'),
                'gender': 'MALE' if gender == 'M' else 'FEMALE',
                'blood_group': random.choice(cls.BLOOD_GROUPS),
                'joining_date': join_date.strftime('%d-%m-%Y'),
                'designation': random.choice(cls.DESIGNATIONS),
                'department': random.choice(cls.DEPARTMENTS),
                'employment_type': random.choice(['PERMANENT', 'PERMANENT', 'CONTRACT', 'TEMPORARY']),
                'address': f'{random.randint(1, 500)}, {random.choice(["Main Street", "Park Road", "Teachers Colony"])}, {random.choice(cls.CITIES)}',
                'salary': str(random.randint(25000, 80000)),
                'experience_years': str(random.randint(1, 25)),
                'emergency_contact_name': f'{random.choice(cls.FIRST_NAMES_MALE)} {last_name}',
                'emergency_contact_phone': f'9{random.randint(100000000, 999999999)}',
            }
            staff.append(staff_member)
        
        return staff
    
    @classmethod
    def generate_classes(cls):
        """Generate sample class and section data"""
        classes = []
        
        for grade in range(1, 13):
            for section in ['A', 'B', 'C']:
                classes.append({
                    'class_name': f'Class {grade}',
                    'section_name': section,
                    'room_number': f'{grade}{section}',
                    'capacity': '40',
                })
        
        return classes
    
    @classmethod
    def generate_subjects(cls):
        """Generate sample subject data"""
        subjects = []
        
        subject_list = [
            ('Mathematics', 'MATH'),
            ('Science', 'SCI'),
            ('English', 'ENG'),
            ('Hindi', 'HIN'),
            ('Social Studies', 'SS'),
            ('Computer Science', 'CS'),
            ('Physical Education', 'PE'),
            ('Art', 'ART'),
        ]
        
        for grade in range(1, 13):
            for name, code in subject_list:
                subjects.append({
                    'subject_name': name,
                    'subject_code': f'{code}{grade}',
                    'class_name': f'Class {grade}',
                    'credit_hours': str(random.randint(3, 6)),
                    'is_elective': 'false' if name in ['Mathematics', 'Science', 'English', 'Hindi'] else 'true',
                    'max_marks': '100',
                })
        
        return subjects
    
    @classmethod
    def generate_fee_structures(cls):
        """Generate sample fee structure data"""
        fees = []
        
        fee_types = [
            ('Tuition Fee', 'MONTHLY', 3000),
            ('Transport Fee', 'MONTHLY', 1500),
            ('Library Fee', 'YEARLY', 500),
            ('Lab Fee', 'YEARLY', 1000),
            ('Sports Fee', 'YEARLY', 800),
            ('Admission Fee', 'ONE_TIME', 5000),
        ]
        
        for grade in range(1, 13):
            for fee_name, frequency, base_amount in fee_types:
                # Increase amount for higher grades
                amount = base_amount + (grade * 100)
                fees.append({
                    'fee_type': fee_name,
                    'class_name': f'Class {grade}',
                    'amount': str(amount),
                    'frequency': frequency,
                    'due_day': '10',
                    'is_mandatory': 'true' if fee_name in ['Tuition Fee', 'Admission Fee'] else 'false',
                    'description': f'{fee_name} for Class {grade}',
                })
        
        return fees
    
    @classmethod
    def to_csv(cls, data, fieldnames=None):
        """Convert data to CSV string"""
        if not data:
            return ''
        
        if fieldnames is None:
            fieldnames = list(data[0].keys())
        
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(data)
        return output.getvalue()
    
    @classmethod
    def to_excel(cls, data, sheet_name='Data', fieldnames=None):
        """Convert data to Excel workbook bytes"""
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export")
        
        if not data:
            return None
        
        if fieldnames is None:
            fieldnames = list(data[0].keys())
        
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # Header styling
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='E5E7EB', fill_type='solid')
        
        # Write headers
        for col, field in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col, value=field)
            cell.font = header_font
            cell.fill = header_fill
        
        # Write data
        for row_num, row_data in enumerate(data, 2):
            for col, field in enumerate(fieldnames, 1):
                ws.cell(row=row_num, column=col, value=row_data.get(field, ''))
        
        # Return bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()


def generate_sample_files(output_dir=None):
    """Generate all sample files for testing"""
    import os
    
    if output_dir is None:
        output_dir = os.path.dirname(os.path.abspath(__file__))
    
    samples_dir = os.path.join(output_dir, 'sample_files')
    os.makedirs(samples_dir, exist_ok=True)
    
    generator = SampleDataGenerator()
    
    # Generate students
    students = generator.generate_students(50)
    with open(os.path.join(samples_dir, 'sample_students.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        f.write(generator.to_csv(students))
    
    if OPENPYXL_AVAILABLE:
        with open(os.path.join(samples_dir, 'sample_students.xlsx'), 'wb') as f:
            f.write(generator.to_excel(students, 'Students'))
    
    # Generate staff
    staff = generator.generate_staff(20)
    with open(os.path.join(samples_dir, 'sample_staff.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        f.write(generator.to_csv(staff))
    
    # Generate classes
    classes = generator.generate_classes()
    with open(os.path.join(samples_dir, 'sample_classes.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        f.write(generator.to_csv(classes))
    
    # Generate subjects
    subjects = generator.generate_subjects()
    with open(os.path.join(samples_dir, 'sample_subjects.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        f.write(generator.to_csv(subjects))
    
    # Generate fee structures
    fees = generator.generate_fee_structures()
    with open(os.path.join(samples_dir, 'sample_fee_structures.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        f.write(generator.to_csv(fees))
    
    print(f"Sample files generated in: {samples_dir}")
    return samples_dir


if __name__ == '__main__':
    generate_sample_files()
